from __future__ import annotations

from xml.etree import ElementTree
from zipfile import ZipFile

import pytest
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches

from pubify_ppt.anchors import shape_alt_text
from pubify_ppt.cli import build_parser, main

def test_cli_help_includes_planned_commands(capsys: pytest.CaptureFixture[str]) -> None:
    parser = build_parser()

    with pytest.raises(SystemExit) as exc_info:
        parser.parse_args(["--help"])

    assert exc_info.value.code == 0
    output = capsys.readouterr().out
    assert "ppt init <presentation-id>" in output
    assert "ppt <presentation-id> figure <figure-id> addto <slide-number>" in output
    assert "ppt <presentation-id> figure <figure-id> update [--output <path>]" in output
    assert "ppt <presentation-id> table <table-id> addto <slide-number>" in output
    assert "ppt <presentation-id> update [--output <path>]" in output


def test_cli_entrypoint_reports_missing_command(capsys: pytest.CaptureFixture[str]) -> None:
    with pytest.raises(SystemExit) as exc_info:
        main([])

    assert exc_info.value.code == 2
    assert "missing command; run 'ppt --help'" in capsys.readouterr().err


def test_cli_init_workspace_creates_pubify_yaml(
    tmp_path,
    monkeypatch: pytest.MonkeyPatch,
    capsys: pytest.CaptureFixture[str],
) -> None:
    monkeypatch.chdir(tmp_path)

    assert main(["init"]) == 0

    assert (tmp_path / "pubify.yaml").read_text(encoding="utf-8") == (
        "pubify-ppt:\n"
        "  presentations_root: slides\n"
    )
    assert (tmp_path / "slides").is_dir()
    assert not (tmp_path / "slides" / "AGENTS.md").exists()
    assert capsys.readouterr().out.strip() == str(tmp_path)


def test_cli_init_rejects_force(
    tmp_path,
    monkeypatch: pytest.MonkeyPatch,
    capsys: pytest.CaptureFixture[str],
) -> None:
    monkeypatch.chdir(tmp_path)

    with pytest.raises(SystemExit) as exc_info:
        main(["init", "--force"])

    assert exc_info.value.code == 2
    assert "init does not accept --force" in capsys.readouterr().err


def test_cli_init_presentation_creates_scaffold(
    tmp_path,
    monkeypatch: pytest.MonkeyPatch,
    capsys: pytest.CaptureFixture[str],
) -> None:
    monkeypatch.chdir(tmp_path)
    main(["init"])
    capsys.readouterr()

    assert main(["init", "demo"]) == 0

    presentation_root = tmp_path / "slides" / "demo"
    assert (presentation_root / "ppt.yaml").is_file()
    assert (presentation_root / "figures.py").is_file()
    assert (presentation_root / "deck.pptx").is_file()
    assert (presentation_root / "data" / "ppt-artifacts" / "figures").is_dir()
    assert (presentation_root / "data" / "ppt-artifacts" / "backups").is_dir()
    assert capsys.readouterr().out.strip() == str(presentation_root)


def test_cli_list_reports_presentations(
    tmp_path,
    monkeypatch: pytest.MonkeyPatch,
    capsys: pytest.CaptureFixture[str],
) -> None:
    monkeypatch.chdir(tmp_path)
    main(["init"])
    main(["init", "demo"])
    main(["init", "talk"])
    capsys.readouterr()

    assert main(["list"]) == 0

    assert capsys.readouterr().out.splitlines() == ["demo", "talk"]


def test_cli_inventory_and_check_commands(
    tmp_path,
    monkeypatch: pytest.MonkeyPatch,
    capsys: pytest.CaptureFixture[str],
) -> None:
    monkeypatch.chdir(tmp_path)
    main(["init"])
    main(["init", "demo"])
    capsys.readouterr()

    assert main(["demo", "data", "list"]) == 0
    assert capsys.readouterr().out.splitlines() == ["example"]
    assert main(["demo", "figure", "list"]) == 0
    assert capsys.readouterr().out.splitlines() == ["example"]
    assert main(["demo", "stat", "list"]) == 0
    assert capsys.readouterr().out.splitlines() == ["example"]
    assert main(["demo", "table", "list"]) == 0
    assert capsys.readouterr().out.splitlines() == ["example"]
    assert main(["demo", "check"]) == 0
    assert capsys.readouterr().out.strip() == "demo: ok"


def test_cli_figure_update_writes_png_and_deck(
    tmp_path,
    monkeypatch: pytest.MonkeyPatch,
    capsys: pytest.CaptureFixture[str],
) -> None:
    monkeypatch.chdir(tmp_path)
    main(["init"])
    main(["init", "demo"])
    capsys.readouterr()

    assert main(["demo", "figure", "update"]) == 0

    assert capsys.readouterr().out.strip() == "Slide 1: {{fig:example}} = example.png"
    assert (tmp_path / "slides" / "demo" / "data" / "ppt-artifacts" / "figures" / "example.png").is_file()


def test_cli_figure_update_warns_once_for_unavailable_theme_font(
    tmp_path,
    monkeypatch: pytest.MonkeyPatch,
    capsys: pytest.CaptureFixture[str],
) -> None:
    monkeypatch.chdir(tmp_path)
    main(["init"])
    main(["init", "demo"])
    presentation_root = tmp_path / "slides" / "demo"
    _replace_theme_body_font(presentation_root / "deck.pptx", "Missing Theme Font")

    def fake_findfont(*args: object, **kwargs: object) -> str:
        raise ValueError("missing font")

    monkeypatch.setattr("pubify_ppt.figures.font_manager.findfont", fake_findfont)
    capsys.readouterr()

    assert main(["demo", "figure", "update"]) == 0

    captured = capsys.readouterr()
    assert captured.out.strip() == "Slide 1: {{fig:example}} = example.png"
    assert captured.err.splitlines() == [
        "Warning: PowerPoint theme figure font 'Missing Theme Font' was not found by Matplotlib; "
        "using Matplotlib's fallback font."
    ]


def test_cli_stat_update_writes_deck(
    tmp_path,
    monkeypatch: pytest.MonkeyPatch,
    capsys: pytest.CaptureFixture[str],
) -> None:
    monkeypatch.chdir(tmp_path)
    main(["init"])
    main(["init", "demo"])
    capsys.readouterr()

    assert main(["demo", "stat", "update"]) == 0

    assert capsys.readouterr().out.splitlines() == ["Slide 1: {{stat:example.count}} = 3"]


def test_cli_table_update_writes_deck(
    tmp_path,
    monkeypatch: pytest.MonkeyPatch,
    capsys: pytest.CaptureFixture[str],
) -> None:
    monkeypatch.chdir(tmp_path)
    main(["init"])
    main(["init", "demo"])
    capsys.readouterr()

    assert main(["demo", "table", "update"]) == 0

    assert capsys.readouterr().out.splitlines() == ["Slide 1: {{table:example}} = 4 rows x 2 columns"]


def test_cli_full_update_writes_figures_stats_and_tables(
    tmp_path,
    monkeypatch: pytest.MonkeyPatch,
    capsys: pytest.CaptureFixture[str],
) -> None:
    monkeypatch.chdir(tmp_path)
    main(["init"])
    main(["init", "demo"])
    capsys.readouterr()

    assert main(["demo", "update"]) == 0

    assert capsys.readouterr().out.splitlines() == [
        "Slide 1: {{fig:example}} = example.png",
        "Slide 1: {{stat:example.count}} = 3",
        "Slide 1: {{table:example}} = 4 rows x 2 columns",
    ]


def test_cli_addto_inserts_anchors_on_existing_slide(
    tmp_path,
    monkeypatch: pytest.MonkeyPatch,
    capsys: pytest.CaptureFixture[str],
) -> None:
    monkeypatch.chdir(tmp_path)
    main(["init"])
    main(["init", "demo"])
    presentation_root = tmp_path / "slides" / "demo"
    deck = Presentation()
    deck.slides.add_slide(deck.slide_layouts[6])
    deck.save(presentation_root / "deck.pptx")
    capsys.readouterr()

    assert main(["demo", "figure", "example:1", "addto", "1"]) == 0
    assert main(["demo", "stat", "example.count", "addto", "1"]) == 0
    assert main(["demo", "table", "example", "addto", "1"]) == 0

    assert capsys.readouterr().out.splitlines() == [
        "Slide 1: {{fig:example:1}} = example.png",
        "Slide 1: {{stat:example.count}} = 3",
        "Slide 1: {{table:example}} = 4 rows x 2 columns",
    ]
    deck = Presentation(presentation_root / "deck.pptx")
    shapes = list(deck.slides[0].shapes)
    assert "{{fig:example:1}}" in [shape_alt_text(shape) for shape in shapes]
    assert any(shape.shape_type == MSO_SHAPE_TYPE.PICTURE for shape in shapes)
    assert "{{table:example}}" in [shape_alt_text(shape) for shape in shapes]
    assert "{{stat:example.count=3}}" in [shape_alt_text(shape) for shape in shapes]
    assert "3" in "\n".join(getattr(shape, "text", "") for shape in shapes)
    assert any(getattr(shape, "has_table", False) for shape in shapes)
    assert (presentation_root / "data" / "ppt-artifacts" / "figures" / "example.png").is_file()


def test_cli_figure_addto_infers_first_panel_for_multi_panel_figure(
    tmp_path,
    monkeypatch: pytest.MonkeyPatch,
    capsys: pytest.CaptureFixture[str],
) -> None:
    monkeypatch.chdir(tmp_path)
    main(["init"])
    main(["init", "demo"])
    presentation_root = tmp_path / "slides" / "demo"
    (presentation_root / "figures.py").write_text(
        "\n".join(
            [
                "import matplotlib.pyplot as plt",
                "from pubify_data import figure",
                "@figure",
                "def plot_pair(ctx):",
                "    fig1, ax1 = plt.subplots()",
                "    ax1.plot([1, 2], [1, 2])",
                "    fig2, ax2 = plt.subplots()",
                "    ax2.plot([1, 2], [2, 1])",
                "    return [fig1, fig2]",
            ]
        )
        + "\n",
        encoding="utf-8",
    )
    deck = Presentation()
    deck.slides.add_slide(deck.slide_layouts[6])
    deck.save(presentation_root / "deck.pptx")
    capsys.readouterr()

    assert main(["demo", "figure", "pair", "addto", "1"]) == 0

    assert capsys.readouterr().out.splitlines() == ["Slide 1: {{fig:pair:1}} = pair_1.png"]
    deck = Presentation(presentation_root / "deck.pptx")
    shapes = list(deck.slides[0].shapes)
    assert "{{fig:pair:1}}" in [shape_alt_text(shape) for shape in shapes]
    assert any(shape.shape_type == MSO_SHAPE_TYPE.PICTURE for shape in shapes)
    assert (presentation_root / "data" / "ppt-artifacts" / "figures" / "pair_1.png").is_file()


def test_cli_table_addto_updates_existing_and_new_table_anchors(
    tmp_path,
    monkeypatch: pytest.MonkeyPatch,
    capsys: pytest.CaptureFixture[str],
) -> None:
    monkeypatch.chdir(tmp_path)
    main(["init"])
    main(["init", "demo"])
    capsys.readouterr()

    assert main(["demo", "table", "example", "addto", "1"]) == 0

    assert capsys.readouterr().out.splitlines() == [
        "Slide 1: {{table:example}} [1/2] = 4 rows x 2 columns",
        "Slide 1: {{table:example}} [2/2] = 4 rows x 2 columns",
    ]
    deck = Presentation(tmp_path / "slides" / "demo" / "deck.pptx")
    assert sum(1 for shape in deck.slides[0].shapes if getattr(shape, "has_table", False)) == 2


def test_cli_addto_rejects_missing_slide(
    tmp_path,
    monkeypatch: pytest.MonkeyPatch,
    capsys: pytest.CaptureFixture[str],
) -> None:
    monkeypatch.chdir(tmp_path)
    main(["init"])
    main(["init", "demo"])
    capsys.readouterr()

    assert main(["demo", "figure", "example", "addto", "2"]) == 1

    assert "Error: Slide number must be between 1 and 1: 2" in capsys.readouterr().err


def test_cli_full_update_reports_replacements_in_slide_order(
    tmp_path,
    monkeypatch: pytest.MonkeyPatch,
    capsys: pytest.CaptureFixture[str],
) -> None:
    monkeypatch.chdir(tmp_path)
    main(["init"])
    main(["init", "demo"])
    presentation_root = tmp_path / "slides" / "demo"
    deck = Presentation()
    first = deck.slides.add_slide(deck.slide_layouts[6])
    first.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(6.0), Inches(0.5)).text = "{{stat:example.count}}"
    second = deck.slides.add_slide(deck.slide_layouts[6])
    second.shapes.add_shape(
        1,
        Inches(0.5),
        Inches(0.5),
        Inches(3),
        Inches(2),
    ).text = "{{fig:example}}"
    second.shapes.add_textbox(Inches(4.0), Inches(0.5), Inches(3), Inches(1.5)).text = "{{table:example}}"
    deck.save(presentation_root / "deck.pptx")
    capsys.readouterr()

    assert main(["demo", "update"]) == 0

    assert capsys.readouterr().out.splitlines() == [
        "Slide 1: {{stat:example.count}} = 3",
        "Slide 2: {{fig:example}} = example.png",
        "Slide 2: {{table:example}} = 4 rows x 2 columns",
    ]


def test_cli_output_writes_generated_copy(
    tmp_path,
    monkeypatch: pytest.MonkeyPatch,
    capsys: pytest.CaptureFixture[str],
) -> None:
    monkeypatch.chdir(tmp_path)
    main(["init"])
    main(["init", "demo"])
    capsys.readouterr()

    assert main(["demo", "stat", "update", "--output", "copy.pptx"]) == 0

    assert (tmp_path / "copy.pptx").is_file()


def test_cli_runtime_error_omits_usage_banner(
    tmp_path,
    monkeypatch: pytest.MonkeyPatch,
    capsys: pytest.CaptureFixture[str],
) -> None:
    monkeypatch.chdir(tmp_path)
    main(["init"])
    main(["init", "demo"])
    deck_path = tmp_path / "slides" / "demo" / "deck.pptx"
    deck = Presentation(deck_path)
    deck.slides[0].shapes.add_textbox(Inches(0.5), Inches(5.5), Inches(4.0), Inches(0.5)).text = (
        "{{stat:example.missing}}"
    )
    deck.save(deck_path)
    capsys.readouterr()

    assert main(["demo", "stat", "update"]) == 1

    captured = capsys.readouterr()
    assert captured.out == ""
    assert "Error: Missing key 'missing' for stat 'example'" in captured.err
    assert "usage:" not in captured.err


def test_cli_labels_repeated_replacements_with_occurrence_index(
    tmp_path,
    monkeypatch: pytest.MonkeyPatch,
    capsys: pytest.CaptureFixture[str],
) -> None:
    monkeypatch.chdir(tmp_path)
    main(["init"])
    main(["init", "demo"])
    deck_path = tmp_path / "slides" / "demo" / "deck.pptx"
    deck = Presentation(deck_path)
    deck.slides[0].shapes.add_textbox(Inches(0.5), Inches(5.5), Inches(4.0), Inches(0.5)).text = (
        "{{stat:example.count}}"
    )
    deck.save(deck_path)
    capsys.readouterr()

    assert main(["demo", "stat", "update"]) == 0

    assert capsys.readouterr().out.splitlines() == [
        "Slide 1: {{stat:example.count}} [1/2] = 3",
        "Slide 1: {{stat:example.count}} [2/2] = 3",
    ]


def _replace_theme_body_font(deck_path, font_family: str) -> None:
    drawing_namespace = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
    with ZipFile(deck_path) as package:
        entries = {name: package.read(name) for name in package.namelist()}
    theme_name = next(name for name in sorted(entries) if name.startswith("ppt/theme/theme") and name.endswith(".xml"))
    root = ElementTree.fromstring(entries[theme_name])
    latin = root.find(f".//{drawing_namespace}minorFont/{drawing_namespace}latin")
    assert latin is not None
    latin.set("typeface", font_family)
    entries[theme_name] = ElementTree.tostring(root, encoding="utf-8", xml_declaration=True)
    with ZipFile(deck_path, "w") as package:
        for name, payload in entries.items():
            package.writestr(name, payload)
