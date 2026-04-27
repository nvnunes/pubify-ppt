from __future__ import annotations

from pathlib import Path

import pytest
from pptx import Presentation

from pubify_ppt.backups import write_deck
from pubify_ppt.discovery import load_presentation_definition
from pubify_ppt.init import init_presentation_by_id, init_workspace
from pubify_ppt.stats import update_stats_to_output


def test_in_place_update_creates_backup(tmp_path: Path) -> None:
    init_workspace(tmp_path)
    init_presentation_by_id(tmp_path, "demo")
    presentation = load_presentation_definition(tmp_path, "demo")

    update_stats_to_output(presentation)

    backups = sorted((tmp_path / "slides" / "demo" / "data" / "ppt-artifacts" / "backups").glob("deck-*.pptx"))
    assert len(backups) == 1
    assert "{{stat:example.count}}" in _deck_text(backups[0])
    assert "Example count: 3" in _deck_text(tmp_path / "slides" / "demo" / "deck.pptx")


def test_output_update_does_not_mutate_source_or_create_backup(tmp_path: Path) -> None:
    init_workspace(tmp_path)
    init_presentation_by_id(tmp_path, "demo")
    presentation = load_presentation_definition(tmp_path, "demo")
    output_path = tmp_path / "copy.pptx"

    update_stats_to_output(presentation, output=output_path)

    assert "{{stat:example.count}}" in _deck_text(tmp_path / "slides" / "demo" / "deck.pptx")
    assert "Example count: 3" in _deck_text(output_path)
    backups = list((tmp_path / "slides" / "demo" / "data" / "ppt-artifacts" / "backups").glob("deck-*.pptx"))
    assert backups == []


def test_in_place_update_raises_when_powerpoint_lock_exists(tmp_path: Path) -> None:
    init_workspace(tmp_path)
    init_presentation_by_id(tmp_path, "demo")
    presentation = load_presentation_definition(tmp_path, "demo")
    _write_powerpoint_lock(presentation.paths.deck_path)

    with pytest.raises(RuntimeError, match="PowerPoint lock file found"):
        update_stats_to_output(presentation)

    assert "{{stat:example.count}}" in _deck_text(presentation.paths.deck_path)


def test_output_update_allows_source_powerpoint_lock(tmp_path: Path) -> None:
    init_workspace(tmp_path)
    init_presentation_by_id(tmp_path, "demo")
    presentation = load_presentation_definition(tmp_path, "demo")
    output_path = tmp_path / "copy.pptx"
    _write_powerpoint_lock(presentation.paths.deck_path)

    update_stats_to_output(presentation, output=output_path)

    assert "{{stat:example.count}}" in _deck_text(presentation.paths.deck_path)
    assert "Example count: 3" in _deck_text(output_path)


def test_lock_error_happens_before_backup_or_temp_source_write(tmp_path: Path) -> None:
    init_workspace(tmp_path)
    init_presentation_by_id(tmp_path, "demo")
    presentation = load_presentation_definition(tmp_path, "demo")
    original_bytes = presentation.paths.deck_path.read_bytes()
    _write_powerpoint_lock(presentation.paths.deck_path)

    with pytest.raises(RuntimeError, match="use --output"):
        update_stats_to_output(presentation)

    assert presentation.paths.deck_path.read_bytes() == original_bytes
    assert list(presentation.paths.backups_root.glob("deck-*.pptx")) == []
    assert list(presentation.paths.deck_path.parent.glob(".deck-*.pptx")) == []


def test_backup_retention_prunes_old_backups(tmp_path: Path) -> None:
    init_workspace(tmp_path)
    init_presentation_by_id(tmp_path, "demo")
    presentation_root = tmp_path / "slides" / "demo"
    (presentation_root / "ppt.yaml").write_text(
        "\n".join(
            [
                "deck: deck.pptx",
                "backup_retention: 2",
                "defaults:",
                "  dpi: 200",
                "external_data_roots:",
            ]
        )
        + "\n",
        encoding="utf-8",
    )

    for index in range(4):
        deck = Presentation(presentation_root / "deck.pptx")
        deck.slides[0].shapes.add_textbox(0, 0, 1000, 1000).text = f"marker {index}"
        deck.save(presentation_root / "deck.pptx")
        presentation = load_presentation_definition(tmp_path, "demo")
        update_stats_to_output(presentation)

    backups = sorted((presentation_root / "data" / "ppt-artifacts" / "backups").glob("deck-*.pptx"))
    assert len(backups) == 2
    assert any("marker 2" in _deck_text(path) for path in backups)
    assert any("marker 3" in _deck_text(path) for path in backups)


def test_failed_in_place_write_keeps_original_and_created_backup(tmp_path: Path) -> None:
    init_workspace(tmp_path)
    init_presentation_by_id(tmp_path, "demo")
    presentation = load_presentation_definition(tmp_path, "demo")
    original_bytes = presentation.paths.deck_path.read_bytes()

    class BrokenDeck:
        def save(self, path: object) -> None:
            raise RuntimeError("write failed")

    try:
        write_deck(presentation, BrokenDeck())
    except RuntimeError:
        pass

    assert presentation.paths.deck_path.read_bytes() == original_bytes
    backups = sorted(presentation.paths.backups_root.glob("deck-*.pptx"))
    assert len(backups) == 1
    assert backups[0].read_bytes() == original_bytes


def _deck_text(deck_path: Path) -> str:
    deck = Presentation(deck_path)
    return "\n".join(shape.text for slide in deck.slides for shape in slide.shapes if hasattr(shape, "text"))


def _write_powerpoint_lock(deck_path: Path) -> None:
    deck_path.with_name(f"~${deck_path.name}").write_text("lock", encoding="utf-8")
