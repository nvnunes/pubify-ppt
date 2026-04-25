from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches
import pytest

from pubify_ppt.anchors import shape_alt_text
from pubify_ppt.discovery import load_presentation_definition
from pubify_ppt.init import init_presentation_by_id, init_workspace
from pubify_ppt.tables import update_tables, update_tables_to_output


def test_update_tables_converts_visible_token_to_native_table(tmp_path: Path) -> None:
    init_workspace(tmp_path)
    init_presentation_by_id(tmp_path, "demo")
    presentation_root = tmp_path / "slides" / "demo"
    _write_table_module(presentation_root, [["A", 1], ["B", 2]], columns=("name", "value"))
    _write_deck_with_visible_table_token(presentation_root, "{{table:summary}}")
    presentation = load_presentation_definition(tmp_path, "demo")

    replacements = update_tables(presentation)

    assert [(item.token, item.summary) for item in replacements] == [("{{table:summary}}", "3 rows x 2 columns")]
    table_shape = _single_table_shape(presentation_root / "deck.pptx")
    assert shape_alt_text(table_shape) == "{{table:summary}}"
    assert _table_text(table_shape) == [["name", "value"], ["A", "1"], ["B", "2"]]


def test_update_tables_converts_alt_text_token_to_native_table(tmp_path: Path) -> None:
    init_workspace(tmp_path)
    init_presentation_by_id(tmp_path, "demo")
    presentation_root = tmp_path / "slides" / "demo"
    _write_table_module(presentation_root, [["A", None]], columns=("name", "value"))
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.5), Inches(0.5), Inches(4), Inches(1.5))
    _set_shape_alt_text(shape, "{{table:summary}}")
    deck.save(presentation_root / "deck.pptx")
    presentation = load_presentation_definition(tmp_path, "demo")

    update_tables(presentation)

    table_shape = _single_table_shape(presentation_root / "deck.pptx")
    assert _table_text(table_shape) == [["name", "value"], ["A", ""]]


def test_update_tables_refreshes_existing_body_cells_and_preserves_headings(tmp_path: Path) -> None:
    init_workspace(tmp_path)
    init_presentation_by_id(tmp_path, "demo")
    presentation_root = tmp_path / "slides" / "demo"
    figures_path = presentation_root / "figures.py"
    figures_path.write_text(_table_module([["A", 1]], columns=("name", "value")), encoding="utf-8")
    _write_deck_with_visible_table_token(presentation_root, "{{table:summary}}")
    presentation = load_presentation_definition(tmp_path, "demo")
    update_tables(presentation)

    deck = Presentation(presentation_root / "deck.pptx")
    table = _single_table_shape_in_deck(deck).table
    table.cell(0, 0).text = "Custom label"
    table.cell(0, 1).text = "Custom total"
    deck.save(presentation_root / "deck.pptx")
    figures_path.write_text(_table_module([["B", 2]], columns=("ignored",)), encoding="utf-8")
    presentation = load_presentation_definition(tmp_path, "demo")
    replacements = update_tables(presentation)

    assert [(item.token, item.summary) for item in replacements] == [("{{table:summary}}", "2 rows x 2 columns")]
    table_shape = _single_table_shape(presentation_root / "deck.pptx")
    assert _table_text(table_shape) == [["Custom label", "Custom total"], ["B", "2"]]


def test_update_tables_reports_column_count_mismatch(tmp_path: Path) -> None:
    init_workspace(tmp_path)
    init_presentation_by_id(tmp_path, "demo")
    presentation_root = tmp_path / "slides" / "demo"
    figures_path = presentation_root / "figures.py"
    figures_path.write_text(_table_module([["A", 1]], columns=("name", "value")), encoding="utf-8")
    _write_deck_with_visible_table_token(presentation_root, "{{table:summary}}")
    presentation = load_presentation_definition(tmp_path, "demo")
    update_tables(presentation)

    figures_path.write_text(_table_module([["A", 1, "x"]], columns=("name", "value", "extra")), encoding="utf-8")
    presentation = load_presentation_definition(tmp_path, "demo")

    with pytest.raises(ValueError, match=r"table \{\{table:summary\}\} is 2x2 but data is 2x3"):
        update_tables(presentation)


def test_update_tables_reports_row_count_mismatch(tmp_path: Path) -> None:
    init_workspace(tmp_path)
    init_presentation_by_id(tmp_path, "demo")
    presentation_root = tmp_path / "slides" / "demo"
    figures_path = presentation_root / "figures.py"
    figures_path.write_text(_table_module([["A", 1]], columns=("name", "value")), encoding="utf-8")
    _write_deck_with_visible_table_token(presentation_root, "{{table:summary}}")
    presentation = load_presentation_definition(tmp_path, "demo")
    update_tables(presentation)

    figures_path.write_text(_table_module([["A", 1], ["B", 2]], columns=("name", "value")), encoding="utf-8")
    presentation = load_presentation_definition(tmp_path, "demo")

    with pytest.raises(ValueError, match="Adjust the table size or replace it with"):
        update_tables(presentation)


def test_update_tables_uses_blank_default_headings_when_columns_metadata_is_missing(tmp_path: Path) -> None:
    init_workspace(tmp_path)
    init_presentation_by_id(tmp_path, "demo")
    presentation_root = tmp_path / "slides" / "demo"
    _write_table_module(presentation_root, [["A", 1]], columns=None)
    _write_deck_with_visible_table_token(presentation_root, "{{table:summary}}")
    presentation = load_presentation_definition(tmp_path, "demo")

    update_tables(presentation)

    assert _table_text(_single_table_shape(presentation_root / "deck.pptx")) == [["", ""], ["A", "1"]]


def test_update_tables_requires_columns_metadata_width_to_match_data(tmp_path: Path) -> None:
    init_workspace(tmp_path)
    init_presentation_by_id(tmp_path, "demo")
    presentation_root = tmp_path / "slides" / "demo"
    _write_table_module(presentation_root, [["A", 1]], columns=("name",))
    _write_deck_with_visible_table_token(presentation_root, "{{table:summary}}")
    presentation = load_presentation_definition(tmp_path, "demo")

    with pytest.raises(ValueError, match="has 1 columns but data has 2"):
        update_tables(presentation)


def test_update_tables_rejects_multi_body_results(tmp_path: Path) -> None:
    init_workspace(tmp_path)
    init_presentation_by_id(tmp_path, "demo")
    presentation_root = tmp_path / "slides" / "demo"
    (presentation_root / "figures.py").write_text(
        "\n".join(
            [
                "from pubify_data import table",
                "from pubify_ppt import TableResult",
                "@table",
                "def tabulate_summary(ctx):",
                "    return TableResult([[[1]], [[2]]], metadata={'columns': ('value',)})",
            ]
        )
        + "\n",
        encoding="utf-8",
    )
    _write_deck_with_visible_table_token(presentation_root, "{{table:summary}}")
    presentation = load_presentation_definition(tmp_path, "demo")

    with pytest.raises(ValueError, match="support one body in v1"):
        update_tables(presentation)


def test_update_tables_targeted_update_leaves_other_tokens_unchanged(tmp_path: Path) -> None:
    init_workspace(tmp_path)
    init_presentation_by_id(tmp_path, "demo")
    presentation_root = tmp_path / "slides" / "demo"
    (presentation_root / "figures.py").write_text(
        "\n".join(
            [
                "from pubify_data import table",
                "from pubify_ppt import TableResult",
                "@table",
                "def tabulate_first(ctx):",
                "    return TableResult([[1]], metadata={'columns': ('one',)})",
                "@table",
                "def tabulate_second(ctx):",
                "    return TableResult([[2]], metadata={'columns': ('two',)})",
            ]
        )
        + "\n",
        encoding="utf-8",
    )
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(2), Inches(0.5)).text = "{{table:first}}"
    slide.shapes.add_textbox(Inches(3), Inches(0.5), Inches(2), Inches(0.5)).text = "{{table:second}}"
    deck.save(presentation_root / "deck.pptx")
    presentation = load_presentation_definition(tmp_path, "demo")

    replacements = update_tables(presentation, table_id="first")

    assert [item.token for item in replacements] == ["{{table:first}}"]
    deck = Presentation(presentation_root / "deck.pptx")
    assert sum(1 for shape in deck.slides[0].shapes if getattr(shape, "has_table", False)) == 1
    assert "{{table:second}}" in _deck_text(presentation_root / "deck.pptx")


def test_update_tables_output_copy_does_not_mutate_source_deck(tmp_path: Path) -> None:
    init_workspace(tmp_path)
    init_presentation_by_id(tmp_path, "demo")
    presentation_root = tmp_path / "slides" / "demo"
    _write_table_module(presentation_root, [["A", 1]], columns=("name", "value"))
    _write_deck_with_visible_table_token(presentation_root, "{{table:summary}}")
    presentation = load_presentation_definition(tmp_path, "demo")

    update_tables_to_output(presentation, output=tmp_path / "copy.pptx")

    assert (tmp_path / "copy.pptx").is_file()
    assert "{{table:summary}}" in _deck_text(presentation_root / "deck.pptx")
    assert _table_text(_single_table_shape(tmp_path / "copy.pptx")) == [["name", "value"], ["A", "1"]]


def _write_table_module(presentation_root: Path, data: list[list[object]], *, columns: tuple[str, ...] | None) -> None:
    (presentation_root / "figures.py").write_text(_table_module(data, columns=columns), encoding="utf-8")


def _table_module(data: list[list[object]], *, columns: tuple[str, ...] | None) -> str:
    metadata = "{}" if columns is None else f"{{'columns': {columns!r}}}"
    return "\n".join(
        [
            "from pubify_data import table",
            "from pubify_ppt import TableResult",
            "@table",
            "def tabulate_summary(ctx):",
            f"    return TableResult({data!r}, metadata={metadata})",
        ]
    ) + "\n"


def _write_deck_with_visible_table_token(presentation_root: Path, token: str) -> None:
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.5), Inches(0.5), Inches(4), Inches(1.5))
    shape.text = token
    deck.save(presentation_root / "deck.pptx")


def _single_table_shape(deck_path: Path) -> object:
    deck = Presentation(deck_path)
    return _single_table_shape_in_deck(deck)


def _single_table_shape_in_deck(deck: Presentation) -> object:
    table_shapes = [shape for slide in deck.slides for shape in slide.shapes if getattr(shape, "has_table", False)]
    assert len(table_shapes) == 1
    return table_shapes[0]


def _table_text(table_shape: object) -> list[list[str]]:
    table = table_shape.table
    return [
        [table.cell(row_index, column_index).text for column_index in range(len(table.columns))]
        for row_index in range(len(table.rows))
    ]


def _deck_text(deck_path: Path) -> str:
    deck = Presentation(deck_path)
    return "\n".join(shape.text for slide in deck.slides for shape in slide.shapes if hasattr(shape, "text"))


def _set_shape_alt_text(shape: object, value: str) -> None:
    c_nv_pr = shape._element.xpath(".//p:cNvPr")[0]
    c_nv_pr.set("descr", value)
