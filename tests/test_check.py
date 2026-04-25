from __future__ import annotations

from pathlib import Path

import pytest
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches

from pubify_ppt.discovery import load_presentation_definition
from pubify_ppt.init import init_presentation_by_id, init_workspace
from pubify_ppt.runtime import check_presentation, validate_presentation_definition


def test_check_presentation_accepts_starter_deck(tmp_path: Path) -> None:
    init_workspace(tmp_path)
    init_presentation_by_id(tmp_path, "demo")

    presentation = load_presentation_definition(tmp_path, "demo")

    check_presentation(presentation)


def test_validate_presentation_reports_missing_local_data(tmp_path: Path) -> None:
    init_workspace(tmp_path)
    init_presentation_by_id(tmp_path, "demo")
    (tmp_path / "slides" / "demo" / "data" / "example.csv").unlink()
    presentation = load_presentation_definition(tmp_path, "demo")

    errors = validate_presentation_definition(presentation)

    assert any("Missing data path for loader 'example'" in error for error in errors)


def test_validate_presentation_reports_unknown_figure_anchor(tmp_path: Path) -> None:
    init_workspace(tmp_path)
    init_presentation_by_id(tmp_path, "demo")
    deck_path = tmp_path / "slides" / "demo" / "deck.pptx"
    deck = Presentation(deck_path)
    shape = deck.slides[0].shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0.5),
        Inches(0.5),
        Inches(1.0),
        Inches(1.0),
    )
    _set_shape_alt_text(shape, "{{fig:missing}}")
    deck.save(deck_path)
    presentation = load_presentation_definition(tmp_path, "demo")

    errors = validate_presentation_definition(presentation)

    assert "Slide 1: unknown figure anchor {{fig:missing}}" in errors


def test_validate_presentation_reports_conflicting_visible_figure_anchor(tmp_path: Path) -> None:
    init_workspace(tmp_path)
    init_presentation_by_id(tmp_path, "demo")
    deck_path = tmp_path / "slides" / "demo" / "deck.pptx"
    deck = Presentation(deck_path)
    shape = deck.slides[0].shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0.5),
        Inches(0.5),
        Inches(1.0),
        Inches(1.0),
    )
    shape.text = "{{fig:other}}"
    _set_shape_alt_text(shape, "{{fig:example}}")
    deck.save(deck_path)
    presentation = load_presentation_definition(tmp_path, "demo")

    errors = validate_presentation_definition(presentation)

    assert (
        "Slide 1: figure anchor alt text '{{fig:example}}' conflicts with "
        "visible figure token '{{fig:other}}'"
    ) in errors


def test_validate_presentation_reports_malformed_visible_figure_anchor(tmp_path: Path) -> None:
    init_workspace(tmp_path)
    init_presentation_by_id(tmp_path, "demo")
    deck_path = tmp_path / "slides" / "demo" / "deck.pptx"
    deck = Presentation(deck_path)
    shape = deck.slides[0].shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0.5),
        Inches(0.5),
        Inches(1.0),
        Inches(1.0),
    )
    shape.text = "{{fig:example:0}}"
    deck.save(deck_path)
    presentation = load_presentation_definition(tmp_path, "demo")

    errors = validate_presentation_definition(presentation)

    assert "Slide 1: malformed figure anchor token '{{fig:example:0}}'" in errors


def test_validate_presentation_reports_unknown_table_anchor(tmp_path: Path) -> None:
    init_workspace(tmp_path)
    init_presentation_by_id(tmp_path, "demo")
    deck_path = tmp_path / "slides" / "demo" / "deck.pptx"
    deck = Presentation(deck_path)
    shape = deck.slides[0].shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0.5),
        Inches(0.5),
        Inches(1.0),
        Inches(1.0),
    )
    shape.text = "{{table:missing}}"
    deck.save(deck_path)
    presentation = load_presentation_definition(tmp_path, "demo")

    errors = validate_presentation_definition(presentation)

    assert "Slide 1: unknown table anchor {{table:missing}}" in errors


def test_validate_presentation_reports_duplicate_table_anchor(tmp_path: Path) -> None:
    init_workspace(tmp_path)
    init_presentation_by_id(tmp_path, "demo")
    deck_path = tmp_path / "slides" / "demo" / "deck.pptx"
    deck = Presentation(deck_path)
    slide = deck.slides[0]
    slide.shapes.add_textbox(Inches(0.5), Inches(5.5), Inches(2.0), Inches(0.5)).text = "{{table:example}}"
    slide.shapes.add_textbox(Inches(3.0), Inches(5.5), Inches(2.0), Inches(0.5)).text = "{{table:example}}"
    deck.save(deck_path)
    presentation = load_presentation_definition(tmp_path, "demo")

    errors = validate_presentation_definition(presentation)

    assert "Slide 1: duplicate table anchor {{table:example}}" in errors


def test_validate_presentation_reports_malformed_visible_table_anchor(tmp_path: Path) -> None:
    init_workspace(tmp_path)
    init_presentation_by_id(tmp_path, "demo")
    deck_path = tmp_path / "slides" / "demo" / "deck.pptx"
    deck = Presentation(deck_path)
    deck.slides[0].shapes.add_textbox(Inches(0.5), Inches(5.5), Inches(2.0), Inches(0.5)).text = "{{table:}}"
    deck.save(deck_path)
    presentation = load_presentation_definition(tmp_path, "demo")

    errors = validate_presentation_definition(presentation)

    assert "Slide 1: malformed table anchor token '{{table:}}'" in errors


def test_validate_presentation_reports_malformed_table_anchor_with_figure_alt_text(tmp_path: Path) -> None:
    init_workspace(tmp_path)
    init_presentation_by_id(tmp_path, "demo")
    deck_path = tmp_path / "slides" / "demo" / "deck.pptx"
    deck = Presentation(deck_path)
    shape = deck.slides[0].shapes.add_textbox(Inches(0.5), Inches(5.5), Inches(2.0), Inches(0.5))
    shape.text = "{{table:}}"
    _set_shape_alt_text(shape, "{{fig:example}}")
    deck.save(deck_path)
    presentation = load_presentation_definition(tmp_path, "demo")

    errors = validate_presentation_definition(presentation)

    assert "Slide 1: malformed table anchor token '{{table:}}'" in errors


def test_check_presentation_accepts_workspace_relative_external_data_root(tmp_path: Path) -> None:
    init_workspace(tmp_path)
    init_presentation_by_id(tmp_path, "demo")
    external_root = tmp_path / "raw-data"
    external_root.mkdir()
    (external_root / "source.txt").write_text("external", encoding="utf-8")
    presentation_root = tmp_path / "slides" / "demo"
    (presentation_root / "ppt.yaml").write_text(
        "\n".join(
            [
                "deck: deck.pptx",
                "backup_retention: 5",
                "defaults:",
                "  image_format: png",
                "  dpi: 200",
                "  fit: contain",
                "external_data_roots:",
                "  raw: raw-data",
            ]
        )
        + "\n",
        encoding="utf-8",
    )
    (presentation_root / "figures.py").write_text(
        "\n".join(
            [
                "from pubify_data import external_data, figure, stat, table",
                "@external_data('raw', 'source.txt')",
                "def load_example(ctx, path):",
                "    return path.read_text(encoding='utf-8')",
                "@figure",
                "def plot_example(ctx, example):",
                "    return example",
                "@stat",
                "def compute_example(ctx, example):",
                "    return {'count': len(example)}",
                "@table",
                "def tabulate_example(ctx, example):",
                "    return [[example]]",
            ]
        )
        + "\n",
        encoding="utf-8",
    )

    presentation = load_presentation_definition(tmp_path, "demo")

    check_presentation(presentation)


def test_check_presentation_accepts_split_run_stat_token(tmp_path: Path) -> None:
    init_workspace(tmp_path)
    init_presentation_by_id(tmp_path, "demo")
    deck_path = tmp_path / "slides" / "demo" / "deck.pptx"
    deck = Presentation(deck_path)
    text_box = deck.slides[0].shapes.add_textbox(Inches(0.5), Inches(5.5), Inches(5.0), Inches(0.5))
    paragraph = text_box.text_frame.paragraphs[0]
    paragraph.add_run().text = "{{stat:example"
    paragraph.add_run().text = ".count}}"
    deck.save(deck_path)
    presentation = load_presentation_definition(tmp_path, "demo")

    check_presentation(presentation)


def test_check_presentation_reports_malformed_split_run_stat_token(tmp_path: Path) -> None:
    init_workspace(tmp_path)
    init_presentation_by_id(tmp_path, "demo")
    deck_path = tmp_path / "slides" / "demo" / "deck.pptx"
    deck = Presentation(deck_path)
    text_box = deck.slides[0].shapes.add_textbox(Inches(0.5), Inches(5.5), Inches(5.0), Inches(0.5))
    paragraph = text_box.text_frame.paragraphs[0]
    paragraph.add_run().text = "{{stat:example"
    paragraph.add_run().text = " count}}"
    deck.save(deck_path)
    presentation = load_presentation_definition(tmp_path, "demo")

    with pytest.raises(ValueError, match="malformed stat token"):
        check_presentation(presentation)


def _set_shape_alt_text(shape: object, value: str) -> None:
    c_nv_pr = shape._element.xpath(".//p:cNvPr")[0]
    c_nv_pr.set("descr", value)
