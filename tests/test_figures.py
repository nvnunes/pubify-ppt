from __future__ import annotations

from pathlib import Path
from xml.etree import ElementTree
from zipfile import ZipFile

import matplotlib.pyplot as plt
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches
import pytest

from pubify_ppt.discovery import load_presentation_definition
from pubify_ppt.figures import _ExportPadding
from pubify_ppt.figures import _asymmetric_tight_bbox
from pubify_ppt.figures import _available_matplotlib_font_family
from pubify_ppt.figures import _resolved_figure_font_family
from pubify_ppt.figures import _render_panel_png
from pubify_ppt.figures import _theme_body_font_family
from pubify_ppt.figures import update_figures
from pubify_ppt.init import init_presentation_by_id, init_workspace


def test_update_figures_replaces_starter_anchor_with_picture(tmp_path: Path) -> None:
    init_workspace(tmp_path)
    init_presentation_by_id(tmp_path, "demo")
    presentation = load_presentation_definition(tmp_path, "demo")

    outputs = update_figures(presentation)

    assert outputs == (tmp_path / "slides" / "demo" / "data" / "ppt-artifacts" / "figures" / "example.png",)
    assert outputs[0].is_file()
    deck = Presentation(tmp_path / "slides" / "demo" / "deck.pptx")
    pictures = [shape for shape in deck.slides[0].shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE]
    assert len(pictures) == 1
    assert _shape_alt_text(pictures[0]) == "{{fig:example}}"


def test_update_figures_bootstraps_anchor_from_exact_shape_text(tmp_path: Path) -> None:
    init_workspace(tmp_path)
    init_presentation_by_id(tmp_path, "demo")
    presentation_root = tmp_path / "slides" / "demo"
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.5), Inches(0.5), Inches(3), Inches(2))
    shape.text = "{{fig:example}}"
    deck.save(presentation_root / "deck.pptx")
    presentation = load_presentation_definition(tmp_path, "demo")

    outputs = update_figures(presentation)

    assert [path.name for path in outputs] == ["example.png"]
    deck = Presentation(presentation_root / "deck.pptx")
    pictures = [shape for shape in deck.slides[0].shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE]
    assert len(pictures) == 1
    assert _shape_alt_text(pictures[0]) == "{{fig:example}}"


def test_update_figures_patches_only_figure_slide_rels_and_media_parts(tmp_path: Path) -> None:
    init_workspace(tmp_path)
    init_presentation_by_id(tmp_path, "demo")
    presentation_root = tmp_path / "slides" / "demo"
    _write_single_line_figure_module(presentation_root, [1, 2, 3])
    deck = Presentation()
    first = deck.slides.add_slide(deck.slide_layouts[6])
    shape = first.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.5), Inches(0.5), Inches(3), Inches(2))
    _set_shape_alt_text(shape, "{{fig:line}}")
    second = deck.slides.add_slide(deck.slide_layouts[6])
    second.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(2), Inches(0.5)).text = "untouched"
    deck.save(presentation_root / "deck.pptx")
    before = _package_payloads(presentation_root / "deck.pptx")
    presentation = load_presentation_definition(tmp_path, "demo")

    update_figures(presentation, figure_id="line")

    changed = _changed_package_entries(before, _package_payloads(presentation_root / "deck.pptx"))
    assert changed <= {
        "[Content_Types].xml",
        "ppt/slides/slide1.xml",
        "ppt/slides/_rels/slide1.xml.rels",
        "ppt/media/image1.png",
    }
    assert "ppt/slides/slide2.xml" not in changed


def test_update_figures_can_update_existing_picture_anchor(tmp_path: Path) -> None:
    init_workspace(tmp_path)
    init_presentation_by_id(tmp_path, "demo")
    presentation = load_presentation_definition(tmp_path, "demo")

    update_figures(presentation)
    presentation = load_presentation_definition(tmp_path, "demo")
    outputs = update_figures(presentation)

    assert outputs[0].name == "example.png"
    deck = Presentation(tmp_path / "slides" / "demo" / "deck.pptx")
    pictures = [shape for shape in deck.slides[0].shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE]
    assert len(pictures) == 1
    assert _shape_alt_text(pictures[0]) == "{{fig:example}}"


def test_theme_body_font_family_reads_powerpoint_minor_font(tmp_path: Path) -> None:
    init_workspace(tmp_path)
    init_presentation_by_id(tmp_path, "demo")
    deck_path = tmp_path / "slides" / "demo" / "deck.pptx"
    _replace_theme_body_font(deck_path, "Deck Body")

    assert _theme_body_font_family(deck_path) == "Deck Body"


def test_available_matplotlib_font_family_warns_once_for_missing_font(
    monkeypatch: pytest.MonkeyPatch,
    capsys: pytest.CaptureFixture[str],
) -> None:
    def fake_findfont(*args: object, **kwargs: object) -> str:
        raise ValueError("missing font")

    monkeypatch.setattr("pubify_ppt.figures.font_manager.findfont", fake_findfont)

    assert _available_matplotlib_font_family("IBM Plex Sans", source="PowerPoint theme") is None
    assert capsys.readouterr().err == (
        "Warning: PowerPoint theme figure font 'IBM Plex Sans' was not found by Matplotlib; "
        "using Matplotlib's fallback font.\n"
    )


def test_resolved_figure_font_family_prefers_ppt_yaml_default(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    init_workspace(tmp_path)
    init_presentation_by_id(tmp_path, "demo")
    presentation_root = tmp_path / "slides" / "demo"
    config_path = presentation_root / "ppt.yaml"
    config_path.write_text(
        config_path.read_text(encoding="utf-8").replace("  dpi: 200\n", "  dpi: 200\n  figure_font_family: Config Body\n"),
        encoding="utf-8",
    )
    _replace_theme_body_font(presentation_root / "deck.pptx", "Theme Body")
    monkeypatch.setattr("pubify_ppt.figures._available_matplotlib_font_family", lambda font, *, source: font)
    presentation = load_presentation_definition(tmp_path, "demo")

    assert _resolved_figure_font_family(presentation) == "Config Body"


def test_resolved_figure_font_family_ignores_unavailable_theme_font(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    init_workspace(tmp_path)
    init_presentation_by_id(tmp_path, "demo")
    presentation_root = tmp_path / "slides" / "demo"
    _replace_theme_body_font(presentation_root / "deck.pptx", "IBM Plex Sans")
    monkeypatch.setattr("pubify_ppt.figures._available_matplotlib_font_family", lambda font, *, source: None)
    presentation = load_presentation_definition(tmp_path, "demo")

    assert _resolved_figure_font_family(presentation) is None


def test_update_figures_passes_resolved_theme_font_family(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    init_workspace(tmp_path)
    init_presentation_by_id(tmp_path, "demo")
    presentation_root = tmp_path / "slides" / "demo"
    _replace_theme_body_font(presentation_root / "deck.pptx", "Theme Body")
    observed = {}

    def fake_save_on_anchor_canvas(
        payload: object,
        output_path: Path,
        *,
        width: float,
        height: float,
        dpi: int,
        font_family: str | None,
        padding: _ExportPadding,
        render_options: dict[str, object],
    ) -> None:
        observed["font_family"] = font_family
        Image.new("RGB", (20, 10), "white").save(output_path)

    monkeypatch.setattr("pubify_ppt.figures._available_matplotlib_font_family", lambda font, *, source: font)
    monkeypatch.setattr("pubify_ppt.figures._save_fig_on_anchor_canvas", fake_save_on_anchor_canvas)
    presentation = load_presentation_definition(tmp_path, "demo")

    update_figures(presentation)

    assert observed["font_family"] == "Theme Body"


def test_update_figures_passes_default_figure_font_sizes(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    init_workspace(tmp_path)
    init_presentation_by_id(tmp_path, "demo")
    presentation_root = tmp_path / "slides" / "demo"
    config_path = presentation_root / "ppt.yaml"
    config_path.write_text(
        config_path.read_text(encoding="utf-8").replace(
            "  dpi: 200\n",
            "  dpi: 200\n"
            "  figure_base_fontsize_pt: 11\n"
            "  figure_axes_labelsize_pt: 11\n"
            "  figure_tick_labelsize_pt: 10\n"
            "  figure_legend_fontsize_pt: 10\n"
            "  figure_title_fontsize_pt: 12\n",
        ),
        encoding="utf-8",
    )
    observed = {}

    def fake_save_on_anchor_canvas(
        payload: object,
        output_path: Path,
        *,
        width: float,
        height: float,
        dpi: int,
        font_family: str | None,
        padding: _ExportPadding,
        render_options: dict[str, object],
    ) -> None:
        observed["style"] = render_options["style"]
        Image.new("RGB", (20, 10), "white").save(output_path)

    monkeypatch.setattr("pubify_ppt.figures._save_fig_on_anchor_canvas", fake_save_on_anchor_canvas)
    presentation = load_presentation_definition(tmp_path, "demo")

    update_figures(presentation)

    assert observed["style"] == {
        "base_fontsize_pt": 11.0,
        "axes_labelsize_pt": 11.0,
        "tick_labelsize_pt": 10.0,
        "legend_fontsize_pt": 10.0,
        "title_fontsize_pt": 12.0,
    }


def test_update_figures_metadata_style_overrides_default_font_sizes(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    init_workspace(tmp_path)
    init_presentation_by_id(tmp_path, "demo")
    presentation_root = tmp_path / "slides" / "demo"
    config_path = presentation_root / "ppt.yaml"
    config_path.write_text(
        config_path.read_text(encoding="utf-8").replace(
            "  dpi: 200\n",
            "  dpi: 200\n  figure_base_fontsize_pt: 11\n  figure_tick_labelsize_pt: 10\n",
        ),
        encoding="utf-8",
    )
    (presentation_root / "figures.py").write_text(
        "\n".join(
            [
                "import matplotlib.pyplot as plt",
                "from pubify_data import figure, stat, table",
                "from pubify_ppt import FigureResult, StatResult, TableResult",
                "@figure",
                "def plot_example(ctx):",
                "    fig, ax = plt.subplots()",
                "    ax.plot([0, 1], [0, 1])",
                "    return FigureResult(fig, metadata={'style': {'tick_labelsize_pt': 8}})",
                "@stat",
                "def compute_example(ctx):",
                "    return StatResult({'count': 3})",
                "@table",
                "def tabulate_example(ctx):",
                "    return TableResult([[1, 2]], metadata={'columns': ['a', 'b']})",
            ]
        )
        + "\n",
        encoding="utf-8",
    )
    observed = {}

    def fake_save_on_anchor_canvas(
        payload: object,
        output_path: Path,
        *,
        width: float,
        height: float,
        dpi: int,
        font_family: str | None,
        padding: _ExportPadding,
        render_options: dict[str, object],
    ) -> None:
        observed["style"] = render_options["style"]
        Image.new("RGB", (20, 10), "white").save(output_path)

    monkeypatch.setattr("pubify_ppt.figures._save_fig_on_anchor_canvas", fake_save_on_anchor_canvas)
    presentation = load_presentation_definition(tmp_path, "demo")

    update_figures(presentation)

    assert observed["style"] == {"base_fontsize_pt": 11.0, "tick_labelsize_pt": 8}


def test_update_figures_passes_resolved_font_family_to_anchor_canvas_path(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    init_workspace(tmp_path)
    init_presentation_by_id(tmp_path, "demo")
    presentation_root = tmp_path / "slides" / "demo"
    _replace_theme_body_font(presentation_root / "deck.pptx", "Theme Body")
    (presentation_root / "figures.py").write_text(
        "\n".join(
            [
                "import matplotlib.pyplot as plt",
                "from pubify_data import figure, stat, table",
                "from pubify_ppt import FigureResult, StatResult, TableResult",
                "@figure",
                "def plot_example(ctx):",
                "    fig, ax = plt.subplots()",
                "    ax.plot([0, 1], [0, 1])",
                "    return FigureResult(fig, metadata={'export_pad_left_inches': 0.01})",
                "@stat",
                "def compute_example(ctx):",
                "    return StatResult({'count': 3})",
                "@table",
                "def tabulate_example(ctx):",
                "    return TableResult([[1, 2]], metadata={'columns': ['a', 'b']})",
            ]
        )
        + "\n",
        encoding="utf-8",
    )
    observed = {}

    def fake_save_on_anchor_canvas(
        payload: object,
        output_path: Path,
        *,
        width: float,
        height: float,
        dpi: int,
        font_family: str | None,
        padding: _ExportPadding,
        render_options: dict[str, object],
    ) -> None:
        observed["font_family"] = font_family
        Image.new("RGB", (20, 10), "white").save(output_path)

    monkeypatch.setattr("pubify_ppt.figures._available_matplotlib_font_family", lambda font, *, source: font)
    monkeypatch.setattr("pubify_ppt.figures._save_fig_on_anchor_canvas", fake_save_on_anchor_canvas)
    presentation = load_presentation_definition(tmp_path, "demo")

    update_figures(presentation)

    assert observed["font_family"] == "Theme Body"


def test_update_figures_deletes_unreferenced_replaced_media(tmp_path: Path) -> None:
    init_workspace(tmp_path)
    init_presentation_by_id(tmp_path, "demo")
    presentation_root = tmp_path / "slides" / "demo"
    _write_colored_line_figure_module(presentation_root, color="red")
    _write_deck_with_figure_anchor(presentation_root, "{{fig:line}}")
    presentation = load_presentation_definition(tmp_path, "demo")
    update_figures(presentation, figure_id="line")
    first_media = _media_payloads(presentation_root / "deck.pptx")
    before_second_update = _package_payloads(presentation_root / "deck.pptx")

    _write_colored_line_figure_module(presentation_root, color="blue")
    presentation = load_presentation_definition(tmp_path, "demo")
    update_figures(presentation, figure_id="line")

    changed = _changed_package_entries(before_second_update, _package_payloads(presentation_root / "deck.pptx"))
    assert changed <= {"ppt/slides/slide1.xml", "ppt/media/image1.png"}
    assert "ppt/media/image1.png" in changed
    second_media = _media_payloads(presentation_root / "deck.pptx")
    assert len(second_media) == 1
    assert second_media != first_media
    assert _unused_image_relationship_ids(presentation_root / "deck.pptx", slide_number=1) == []
    assert _duplicate_c_nv_pr_ids(presentation_root / "deck.pptx", slide_number=1) == []


def test_update_figures_keeps_existing_picture_geometry_when_export_content_aspect_changes(tmp_path: Path) -> None:
    init_workspace(tmp_path)
    init_presentation_by_id(tmp_path, "demo")
    presentation_root = tmp_path / "slides" / "demo"
    _write_padded_line_figure_module(presentation_root, left_padding=0.0)
    _write_deck_with_figure_anchor(presentation_root, "{{fig:line}}", width=3.4)
    presentation = load_presentation_definition(tmp_path, "demo")
    (first_output,) = update_figures(presentation, figure_id="line")
    first_aspect = _png_aspect(first_output)
    first_geometry = _single_picture_geometry(presentation_root / "deck.pptx")
    before_second_update = _package_payloads(presentation_root / "deck.pptx")

    _write_padded_line_figure_module(presentation_root, left_padding=0.08)
    presentation = load_presentation_definition(tmp_path, "demo")
    (second_output,) = update_figures(presentation, figure_id="line")

    changed = _changed_package_entries(before_second_update, _package_payloads(presentation_root / "deck.pptx"))
    assert changed == {"ppt/media/image1.png"}
    second_geometry = _single_picture_geometry(presentation_root / "deck.pptx")
    assert second_geometry == first_geometry
    assert _png_aspect(second_output) == pytest.approx(first_aspect, rel=0.01)


def test_update_figures_detaches_copied_managed_picture_anchors(tmp_path: Path) -> None:
    init_workspace(tmp_path)
    init_presentation_by_id(tmp_path, "demo")
    presentation_root = tmp_path / "slides" / "demo"
    _write_two_line_figure_module(presentation_root)
    seed = presentation_root / "seed.png"
    Image.new("RGB", (20, 20), "blue").save(seed)
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    first = slide.shapes.add_picture(str(seed), Inches(0.5), Inches(0.5), width=Inches(2))
    second = slide.shapes.add_picture(str(seed), Inches(3), Inches(0.5), width=Inches(2))
    _set_shape_alt_text(first, "{{fig:first}}")
    _set_shape_alt_text(second, "{{fig:second}}")
    deck.save(presentation_root / "deck.pptx")
    assert _picture_blip_ids(presentation_root / "deck.pptx", slide_number=1) == ["rId2", "rId2"]
    presentation = load_presentation_definition(tmp_path, "demo")

    outputs = update_figures(presentation)

    assert [path.name for path in outputs] == ["first.png", "second.png"]
    blip_ids = _picture_blip_ids(presentation_root / "deck.pptx", slide_number=1)
    assert len(set(blip_ids)) == 2
    targets = _image_relationship_targets(presentation_root / "deck.pptx", slide_number=1)
    assert len({targets[r_id] for r_id in blip_ids}) == 2
    media_payloads = _package_payloads(presentation_root / "deck.pptx")
    assert media_payloads[f"ppt/media/{targets[blip_ids[0]].split('/')[-1]}"] == outputs[0].read_bytes()
    assert media_payloads[f"ppt/media/{targets[blip_ids[1]].split('/')[-1]}"] == outputs[1].read_bytes()


def test_update_figures_removes_stale_image_relationships(tmp_path: Path) -> None:
    init_workspace(tmp_path)
    init_presentation_by_id(tmp_path, "demo")
    presentation_root = tmp_path / "slides" / "demo"
    _write_single_line_figure_module(presentation_root, [1, 2, 3])
    _write_deck_with_figure_anchor(presentation_root, "{{fig:line}}")
    presentation = load_presentation_definition(tmp_path, "demo")
    update_figures(presentation, figure_id="line")

    _write_single_line_figure_module(presentation_root, [3, 1, 2])
    presentation = load_presentation_definition(tmp_path, "demo")
    update_figures(presentation, figure_id="line")

    assert _unused_image_relationship_ids(presentation_root / "deck.pptx", slide_number=1) == []


def test_update_figures_supports_targeted_multi_panel_results(tmp_path: Path) -> None:
    init_workspace(tmp_path)
    init_presentation_by_id(tmp_path, "demo")
    presentation_root = tmp_path / "slides" / "demo"
    (presentation_root / "figures.py").write_text(
        "\n".join(
            [
                "import matplotlib.pyplot as plt",
                "from pubify_data import figure",
                "@figure",
                "def plot_pair(ctx):",
                "    fig1, ax1 = plt.subplots()",
                "    ax1.plot([1, 2], [1, 2])",
                "    fig2, ax2 = plt.subplots()",
                "    ax2.plot([1, 2], [2, 1])",
                "    return [fig1, fig2]",
            ]
        )
        + "\n",
        encoding="utf-8",
    )
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    first = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.5), Inches(0.5), Inches(3), Inches(2))
    second = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(4), Inches(0.5), Inches(3), Inches(2))
    _set_shape_alt_text(first, "{{fig:pair:1}}")
    _set_shape_alt_text(second, "{{fig:pair:2}}")
    deck.save(presentation_root / "deck.pptx")
    presentation = load_presentation_definition(tmp_path, "demo")

    outputs = update_figures(presentation, figure_id="pair")

    assert [path.name for path in outputs] == ["pair_1.png", "pair_2.png"]
    deck = Presentation(presentation_root / "deck.pptx")
    pictures = [shape for shape in deck.slides[0].shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE]
    assert sorted(_shape_alt_text(shape) for shape in pictures) == ["{{fig:pair:1}}", "{{fig:pair:2}}"]


def test_update_figures_prepares_axes_payload_without_mutating_source(tmp_path: Path) -> None:
    init_workspace(tmp_path)
    init_presentation_by_id(tmp_path, "demo")
    presentation_root = tmp_path / "slides" / "demo"
    (presentation_root / "figures.py").write_text(
        "\n".join(
            [
                "import matplotlib.pyplot as plt",
                "from pubify_data import figure",
                "from pubify_ppt import FigureResult",
                "@figure",
                "def plot_axes(ctx):",
                "    fig, axs = plt.subplots(1, 2)",
                "    axs[0].set_xlabel('left')",
                "    axs[1].set_xlabel('right')",
                "    axs[0].plot([1, 2], [1, 2])",
                "    axs[1].plot([1, 2], [2, 1])",
                "    return FigureResult(axs[1], metadata={'hide_labels': True})",
            ]
        )
        + "\n",
        encoding="utf-8",
    )
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.5), Inches(0.5), Inches(3), Inches(2))
    _set_shape_alt_text(shape, "{{fig:axes}}")
    deck.save(presentation_root / "deck.pptx")
    presentation = load_presentation_definition(tmp_path, "demo")

    outputs = update_figures(presentation, figure_id="axes")

    assert [path.name for path in outputs] == ["axes.png"]
    assert outputs[0].is_file()


def test_update_figures_saves_reused_axes_panel_with_tight_bbox(tmp_path: Path) -> None:
    init_workspace(tmp_path)
    init_presentation_by_id(tmp_path, "demo")
    presentation_root = tmp_path / "slides" / "demo"
    (presentation_root / "figures.py").write_text(
        "\n".join(
            [
                "import matplotlib.pyplot as plt",
                "from pubify_data import figure",
                "from pubify_ppt import FigureResult",
                "@figure",
                "def plot_reused_axes(ctx):",
                "    fig = plt.figure(figsize=(5, 5))",
                "    ax = fig.add_axes([0.1275, 0.11, 0.77, 0.77])",
                "    ax.plot([0, 1, 2], [0, 1, 0])",
                "    ax.set_xlabel('x label that must remain visible')",
                "    ax.set_ylabel('y label that must remain visible')",
                "    return FigureResult(ax)",
            ]
        )
        + "\n",
        encoding="utf-8",
    )
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.5), Inches(0.5), Inches(3), Inches(2))
    _set_shape_alt_text(shape, "{{fig:reused_axes}}")
    deck.save(presentation_root / "deck.pptx")
    presentation = load_presentation_definition(tmp_path, "demo")

    outputs = update_figures(presentation, figure_id="reused_axes")

    assert [path.name for path in outputs] == ["reused_axes.png"]
    with Image.open(outputs[0]) as image:
        assert image.size == (600, 400)


def test_render_panel_png_expands_layout_to_fill_anchor_canvas(tmp_path: Path) -> None:
    fig, ax = plt.subplots()
    fig.subplots_adjust(left=0.38, right=0.62, bottom=0.38, top=0.62)
    ax.plot([0, 1], [0, 1])
    ax.set_xlabel("x")
    ax.set_ylabel("y")
    output = tmp_path / "expanded.png"

    try:
        _render_panel_png(
            fig,
            output,
            width_emu=Inches(3),
            height_emu=Inches(2),
            dpi=200,
            font_family=None,
            preparation_options={},
        )
    finally:
        plt.close(fig)

    with Image.open(output) as image:
        assert image.size == (600, 400)
    bbox = _non_blank_bbox(output)
    assert (bbox[2] - bbox[0]) / 600 > 0.9
    assert (bbox[3] - bbox[1]) / 400 > 0.9


def test_update_figures_accepts_symmetric_export_padding(tmp_path: Path) -> None:
    init_workspace(tmp_path)
    init_presentation_by_id(tmp_path, "demo")
    presentation_root = tmp_path / "slides" / "demo"
    (presentation_root / "figures.py").write_text(
        "\n".join(
            [
                "import matplotlib.pyplot as plt",
                "from pubify_data import figure",
                "from pubify_ppt import FigureResult",
                "def make_fig():",
                "    fig, ax = plt.subplots()",
                "    ax.plot([0, 1], [0, 1])",
                "    return fig",
                "@figure",
                "def plot_plain(ctx):",
                "    return make_fig()",
                "@figure",
                "def plot_padded(ctx):",
                "    return FigureResult(make_fig(), metadata={'export_pad_inches': 0.05})",
            ]
        )
        + "\n",
        encoding="utf-8",
    )
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    first = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.5), Inches(0.5), Inches(3), Inches(2))
    second = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(4), Inches(0.5), Inches(3), Inches(2))
    _set_shape_alt_text(first, "{{fig:plain}}")
    _set_shape_alt_text(second, "{{fig:padded}}")
    deck.save(presentation_root / "deck.pptx")
    presentation = load_presentation_definition(tmp_path, "demo")

    outputs = update_figures(presentation)

    by_name = {path.name: path for path in outputs}
    with Image.open(by_name["plain.png"]) as plain, Image.open(by_name["padded.png"]) as padded:
        assert plain.size == (600, 400)
        assert padded.size == (600, 400)


def test_update_figures_applies_side_specific_export_padding(tmp_path: Path) -> None:
    init_workspace(tmp_path)
    init_presentation_by_id(tmp_path, "demo")
    presentation_root = tmp_path / "slides" / "demo"
    _write_black_left_padded_figure_module(presentation_root, left_padding=0.0)
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.5), Inches(0.5), Inches(3.2), Inches(2))
    _set_shape_alt_text(shape, "{{fig:left_padded}}")
    deck.save(presentation_root / "deck.pptx")
    presentation = load_presentation_definition(tmp_path, "demo")

    (plain_output,) = update_figures(presentation, figure_id="left_padded")
    plain_first_content_column = _first_non_blank_column(plain_output)
    _write_black_left_padded_figure_module(presentation_root, left_padding=0.08)
    presentation = load_presentation_definition(tmp_path, "demo")
    (output,) = update_figures(presentation, figure_id="left_padded")

    with Image.open(output) as image:
        padding_pixel = image.getpixel((0, image.height // 2))
        assert _is_blank_export_padding_pixel(padding_pixel)
        assert _is_opaque_pixel(padding_pixel)
    assert _first_non_blank_column(output) > plain_first_content_column


def test_side_specific_export_padding_expands_matplotlib_tight_bbox() -> None:
    fig, ax = plt.subplots()
    ax.plot([0, 1], [0, 1])
    try:
        plain = _asymmetric_tight_bbox(
            fig,
            _ExportPadding(
                symmetric=0.0,
                left=0.0,
                right=0.0,
                top=0.0,
                bottom=0.0,
                has_side_specific=True,
            ),
        )
        padded = _asymmetric_tight_bbox(
            fig,
            _ExportPadding(
                symmetric=0.0,
                left=0.06,
                right=0.02,
                top=0.03,
                bottom=0.01,
                has_side_specific=True,
            ),
        )
    finally:
        plt.close(fig)

    assert padded.x0 == pytest.approx(plain.x0 - 0.06)
    assert padded.x1 == pytest.approx(plain.x1 + 0.02)
    assert padded.y0 == pytest.approx(plain.y0 - 0.01)
    assert padded.y1 == pytest.approx(plain.y1 + 0.03)


def test_update_figures_rejects_invalid_export_padding_metadata(tmp_path: Path) -> None:
    init_workspace(tmp_path)
    init_presentation_by_id(tmp_path, "demo")
    presentation_root = tmp_path / "slides" / "demo"
    (presentation_root / "figures.py").write_text(
        "\n".join(
            [
                "import matplotlib.pyplot as plt",
                "from pubify_data import figure",
                "from pubify_ppt import FigureResult",
                "@figure",
                "def plot_bad_padding(ctx):",
                "    fig, ax = plt.subplots()",
                "    ax.plot([1, 2], [1, 2])",
                "    return FigureResult(fig, metadata={'export_pad_left_inches': -0.01})",
            ]
        )
        + "\n",
        encoding="utf-8",
    )
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.5), Inches(0.5), Inches(3), Inches(2))
    _set_shape_alt_text(shape, "{{fig:bad_padding}}")
    deck.save(presentation_root / "deck.pptx")
    presentation = load_presentation_definition(tmp_path, "demo")

    with pytest.raises(ValueError, match="export_pad_left_inches must be a non-negative number"):
        update_figures(presentation, figure_id="bad_padding")


def test_targeted_figure_update_preserves_prefix_matching_png(tmp_path: Path) -> None:
    init_workspace(tmp_path)
    init_presentation_by_id(tmp_path, "demo")
    presentation_root = tmp_path / "slides" / "demo"
    figures_root = presentation_root / "data" / "ppt-artifacts" / "figures"
    figures_root.mkdir(parents=True, exist_ok=True)
    (figures_root / "foo_1.png").write_text("stale panel", encoding="utf-8")
    (figures_root / "foo_bar.png").write_text("separate figure", encoding="utf-8")
    (presentation_root / "figures.py").write_text(
        "\n".join(
            [
                "import matplotlib.pyplot as plt",
                "from pubify_data import figure",
                "@figure",
                "def plot_foo(ctx):",
                "    fig, ax = plt.subplots()",
                "    ax.plot([1, 2], [1, 2])",
                "    return fig",
            ]
        )
        + "\n",
        encoding="utf-8",
    )
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.5), Inches(0.5), Inches(3), Inches(2))
    _set_shape_alt_text(shape, "{{fig:foo}}")
    deck.save(presentation_root / "deck.pptx")
    presentation = load_presentation_definition(tmp_path, "demo")

    outputs = update_figures(presentation, figure_id="foo")

    assert [path.name for path in outputs] == ["foo.png"]
    assert not (figures_root / "foo_1.png").exists()
    assert (figures_root / "foo_bar.png").read_text(encoding="utf-8") == "separate figure"


def test_update_figures_rejects_non_integer_metadata_dpi(tmp_path: Path) -> None:
    init_workspace(tmp_path)
    init_presentation_by_id(tmp_path, "demo")
    presentation_root = tmp_path / "slides" / "demo"
    (presentation_root / "figures.py").write_text(
        "\n".join(
            [
                "import matplotlib.pyplot as plt",
                "from pubify_data import figure",
                "from pubify_ppt import FigureResult",
                "@figure",
                "def plot_bad_dpi(ctx):",
                "    fig, ax = plt.subplots()",
                "    ax.plot([1, 2], [1, 2])",
                "    return FigureResult(fig, metadata={'dpi': '300'})",
            ]
        )
        + "\n",
        encoding="utf-8",
    )
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.5), Inches(0.5), Inches(3), Inches(2))
    _set_shape_alt_text(shape, "{{fig:bad_dpi}}")
    deck.save(presentation_root / "deck.pptx")
    presentation = load_presentation_definition(tmp_path, "demo")

    with pytest.raises(ValueError, match="dpi must be a positive integer"):
        update_figures(presentation, figure_id="bad_dpi")


def test_update_figures_renders_local_wrapper_around_source_publication_figure(tmp_path: Path) -> None:
    init_workspace(tmp_path)
    init_presentation_by_id(tmp_path, "demo")
    source_root = tmp_path / "papers" / "ao4elt8"
    source_data = source_root / "data"
    source_data.mkdir(parents=True)
    (source_data / "points.csv").write_text("x,y\n1,1\n2,4\n", encoding="utf-8")
    (source_root / "figures.py").write_text(
        "\n".join(
            [
                "import csv",
                "import matplotlib.pyplot as plt",
                "from pubify_data import data, figure",
                "@data('points.csv')",
                "def load_points(ctx, path):",
                "    with path.open(newline='', encoding='utf-8') as handle:",
                "        return [(float(row['x']), float(row['y'])) for row in csv.DictReader(handle)]",
                "@figure",
                "def plot_source(ctx, points):",
                "    fig, ax = plt.subplots()",
                "    ax.plot([point[0] for point in points], [point[1] for point in points])",
                "    return fig",
            ]
        )
        + "\n",
        encoding="utf-8",
    )
    presentation_root = tmp_path / "slides" / "demo"
    (presentation_root / "figures.py").write_text(
        "\n".join(
            [
                "from pubify_data import figure",
                "from pubify_ppt import FigureResult",
                "@figure",
                "def plot_source_panel(ctx):",
                "    return FigureResult(ctx.source('ao4elt8').figure('source').panel(1))",
            ]
        )
        + "\n",
        encoding="utf-8",
    )
    (presentation_root / "ppt.yaml").write_text(
        "\n".join(
            [
                "deck: deck.pptx",
                "backup_retention: 5",
                "defaults:",
                "  dpi: 200",
                "external_data_roots:",
                "sources:",
                "  ao4elt8: papers/ao4elt8",
            ]
        )
        + "\n",
        encoding="utf-8",
    )
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.5), Inches(0.5), Inches(3), Inches(2))
    _set_shape_alt_text(shape, "{{fig:source_panel}}")
    deck.save(presentation_root / "deck.pptx")
    presentation = load_presentation_definition(tmp_path, "demo")

    outputs = update_figures(presentation, figure_id="source_panel")

    assert [path.name for path in outputs] == ["source_panel.png"]
    deck = Presentation(presentation_root / "deck.pptx")
    pictures = [shape for shape in deck.slides[0].shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE]
    assert [_shape_alt_text(shape) for shape in pictures] == ["{{fig:source_panel}}"]


def test_update_figures_requires_selected_anchor(tmp_path: Path) -> None:
    init_workspace(tmp_path)
    init_presentation_by_id(tmp_path, "demo")
    deck = Presentation()
    deck.slides.add_slide(deck.slide_layouts[6])
    deck.save(tmp_path / "slides" / "demo" / "deck.pptx")
    presentation = load_presentation_definition(tmp_path, "demo")

    with pytest.raises(ValueError, match="requires exactly one anchor"):
        update_figures(presentation, figure_id="example")


def _shape_alt_text(shape: object) -> str | None:
    c_nv_pr = shape._element.xpath(".//p:cNvPr")[0]
    return c_nv_pr.get("descr")


def _set_shape_alt_text(shape: object, value: str) -> None:
    c_nv_pr = shape._element.xpath(".//p:cNvPr")[0]
    c_nv_pr.set("descr", value)


def _write_deck_with_figure_anchor(
    presentation_root: Path,
    token: str,
    *,
    width: float = 3.0,
    height: float = 2.0,
) -> None:
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.5), Inches(0.5), Inches(width), Inches(height))
    _set_shape_alt_text(shape, token)
    deck.save(presentation_root / "deck.pptx")


def _write_single_line_figure_module(presentation_root: Path, y_values: list[int]) -> None:
    (presentation_root / "figures.py").write_text(
        "\n".join(
            [
                "import matplotlib.pyplot as plt",
                "from pubify_data import figure",
                "@figure",
                "def plot_line(ctx):",
                "    fig, ax = plt.subplots()",
                f"    ax.plot([1, 2, 3], {y_values!r})",
                "    return fig",
            ]
        )
        + "\n",
        encoding="utf-8",
    )


def _write_colored_line_figure_module(presentation_root: Path, *, color: str) -> None:
    (presentation_root / "figures.py").write_text(
        "\n".join(
            [
                "import matplotlib.pyplot as plt",
                "from pubify_data import figure",
                "@figure",
                "def plot_line(ctx):",
                "    fig, ax = plt.subplots()",
                f"    fig.patch.set_facecolor({color!r})",
                "    ax.set_xlim(0, 1)",
                "    ax.set_ylim(0, 1)",
                f"    ax.plot([0, 1], [0, 1], color={color!r})",
                "    return fig",
            ]
        )
        + "\n",
        encoding="utf-8",
    )


def _write_padded_line_figure_module(
    presentation_root: Path,
    *,
    left_padding: float,
    figure_id: str = "line",
) -> None:
    (presentation_root / "pad.txt").write_text(str(left_padding), encoding="utf-8")
    (presentation_root / "figures.py").write_text(
        "\n".join(
            [
                "from pathlib import Path",
                "import matplotlib.pyplot as plt",
                "from pubify_data import figure",
                "from pubify_ppt import FigureResult",
                "@figure",
                f"def plot_{figure_id}(ctx):",
                "    fig, ax = plt.subplots()",
                "    ax.plot([1, 2, 3], [1, 2, 3])",
                "    pad = float(Path(__file__).with_name('pad.txt').read_text(encoding='utf-8'))",
                "    return FigureResult(fig, metadata={'export_pad_left_inches': pad})",
            ]
        )
        + "\n",
        encoding="utf-8",
    )


def _write_black_left_padded_figure_module(presentation_root: Path, *, left_padding: float) -> None:
    (presentation_root / "figures.py").write_text(
        "\n".join(
            [
                "import matplotlib.pyplot as plt",
                "from pubify_data import figure",
                "from pubify_ppt import FigureResult",
                "@figure",
                "def plot_left_padded(ctx):",
                "    fig, ax = plt.subplots()",
                "    ax.set_facecolor('black')",
                "    ax.set_ylabel('left edge label')",
                "    ax.plot([0, 1], [0, 1], color='white')",
                f"    return FigureResult(fig, metadata={{'export_pad_left_inches': {left_padding!r}}})",
            ]
        )
        + "\n",
        encoding="utf-8",
    )


def _write_two_line_figure_module(presentation_root: Path) -> None:
    (presentation_root / "figures.py").write_text(
        "\n".join(
            [
                "import matplotlib.pyplot as plt",
                "from pubify_data import figure",
                "@figure",
                "def plot_first(ctx):",
                "    fig, ax = plt.subplots()",
                "    ax.plot([1, 2, 3], [1, 2, 3])",
                "    return fig",
                "@figure",
                "def plot_second(ctx):",
                "    fig, ax = plt.subplots()",
                "    ax.plot([1, 2, 3], [3, 1, 2])",
                "    return fig",
            ]
        )
        + "\n",
        encoding="utf-8",
    )


def _unused_image_relationship_ids(deck_path: Path, *, slide_number: int) -> list[str]:
    relationship_namespace = "{http://schemas.openxmlformats.org/package/2006/relationships}"
    relationship_attribute_namespace = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
    image_relationship_type = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/image"
    with ZipFile(deck_path) as package:
        slide_xml = package.read(f"ppt/slides/slide{slide_number}.xml")
        rels_xml = package.read(f"ppt/slides/_rels/slide{slide_number}.xml.rels")
    slide = ElementTree.fromstring(slide_xml)
    rels = ElementTree.fromstring(rels_xml)
    used_ids = {
        value
        for element in slide.iter()
        if (value := element.attrib.get(f"{relationship_attribute_namespace}embed")) is not None
    }
    image_ids = {
        relationship.attrib["Id"]
        for relationship in rels.findall(f"{relationship_namespace}Relationship")
        if relationship.attrib.get("Type") == image_relationship_type
    }
    return sorted(image_ids - used_ids)


def _picture_blip_ids(deck_path: Path, *, slide_number: int) -> list[str]:
    relationship_attribute_namespace = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
    with ZipFile(deck_path) as package:
        slide_xml = package.read(f"ppt/slides/slide{slide_number}.xml")
    slide = ElementTree.fromstring(slide_xml)
    return [
        value
        for element in slide.iter()
        if element.tag.endswith("}blip")
        if (value := element.attrib.get(f"{relationship_attribute_namespace}embed")) is not None
    ]


def _image_relationship_targets(deck_path: Path, *, slide_number: int) -> dict[str, str]:
    relationship_namespace = "{http://schemas.openxmlformats.org/package/2006/relationships}"
    image_relationship_type = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/image"
    with ZipFile(deck_path) as package:
        rels_xml = package.read(f"ppt/slides/_rels/slide{slide_number}.xml.rels")
    rels = ElementTree.fromstring(rels_xml)
    return {
        relationship.attrib["Id"]: relationship.attrib["Target"]
        for relationship in rels.findall(f"{relationship_namespace}Relationship")
        if relationship.attrib.get("Type") == image_relationship_type
    }


def _package_payloads(deck_path: Path) -> dict[str, bytes]:
    with ZipFile(deck_path) as package:
        return {name: package.read(name) for name in package.namelist()}


def _changed_package_entries(before: dict[str, bytes], after: dict[str, bytes]) -> set[str]:
    names = set(before) | set(after)
    return {name for name in names if before.get(name) != after.get(name)}


def _media_payloads(deck_path: Path) -> dict[str, bytes]:
    with ZipFile(deck_path) as package:
        return {
            name: package.read(name)
            for name in package.namelist()
            if name.startswith("ppt/media/")
        }


def _replace_theme_body_font(deck_path: Path, font_family: str) -> None:
    drawing_namespace = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
    with ZipFile(deck_path) as package:
        entries = {name: package.read(name) for name in package.namelist()}
    theme_name = next(name for name in sorted(entries) if name.startswith("ppt/theme/theme") and name.endswith(".xml"))
    root = ElementTree.fromstring(entries[theme_name])
    latin = root.find(f".//{drawing_namespace}minorFont/{drawing_namespace}latin")
    assert latin is not None
    latin.set("typeface", font_family)
    entries[theme_name] = ElementTree.tostring(root, encoding="utf-8", xml_declaration=True)
    with ZipFile(deck_path, "w") as package:
        for name, payload in entries.items():
            package.writestr(name, payload)


def _duplicate_c_nv_pr_ids(deck_path: Path, *, slide_number: int) -> list[str]:
    presentation_namespace = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
    with ZipFile(deck_path) as package:
        slide_xml = package.read(f"ppt/slides/slide{slide_number}.xml")
    slide = ElementTree.fromstring(slide_xml)
    ids = [
        element.attrib["id"]
        for element in slide.iter(f"{presentation_namespace}cNvPr")
        if "id" in element.attrib
    ]
    return sorted({value for value in ids if ids.count(value) > 1})


def _single_picture_geometry(deck_path: Path) -> tuple[int, int, int, int]:
    deck = Presentation(deck_path)
    pictures = [shape for slide in deck.slides for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE]
    assert len(pictures) == 1
    picture = pictures[0]
    return int(picture.left), int(picture.top), int(picture.width), int(picture.height)


def _png_aspect(path: Path) -> float:
    with Image.open(path) as image:
        return _aspect_ratio(image.width, image.height)


def _aspect_ratio(width: int, height: int) -> float:
    return width / height


def _first_non_blank_column(path: Path) -> int:
    with Image.open(path) as image:
        for x in range(image.width):
            for y in range(image.height):
                if not _is_blank_export_padding_pixel(image.getpixel((x, y))):
                    return x
    raise AssertionError("image did not contain any non-blank pixels")


def _non_blank_bbox(path: Path) -> tuple[int, int, int, int]:
    with Image.open(path) as image:
        xs = []
        ys = []
        for x in range(image.width):
            for y in range(image.height):
                if not _is_blank_export_padding_pixel(image.getpixel((x, y))):
                    xs.append(x)
                    ys.append(y)
    if not xs:
        raise AssertionError("image did not contain any non-blank pixels")
    return min(xs), min(ys), max(xs) + 1, max(ys) + 1


def _is_blank_export_padding_pixel(pixel: object) -> bool:
    if isinstance(pixel, int):
        return pixel == 255
    channels = tuple(pixel)
    if len(channels) >= 4 and channels[3] == 0:
        return True
    return len(channels) >= 3 and channels[:3] == (255, 255, 255)


def _is_opaque_pixel(pixel: object) -> bool:
    if isinstance(pixel, int):
        return True
    channels = tuple(pixel)
    return len(channels) < 4 or channels[3] == 255
