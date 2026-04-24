from __future__ import annotations

from pathlib import Path

import pytest
from pptx import Presentation

from pubify_ppt.init import init_presentation_by_id, init_workspace


def test_init_workspace_appends_pubify_ppt_section_to_existing_config(tmp_path: Path) -> None:
    (tmp_path / "pubify.yaml").write_text(
        "pubify-pubs:\n  publications_root: papers\n",
        encoding="utf-8",
    )

    init_workspace(tmp_path)

    assert (tmp_path / "pubify.yaml").read_text(encoding="utf-8") == (
        "pubify-pubs:\n"
        "  publications_root: papers\n"
        "\n"
        "pubify-ppt:\n"
        "  presentations_root: slides\n"
    )
    assert (tmp_path / "slides").is_dir()


def test_init_workspace_does_not_duplicate_existing_section(tmp_path: Path) -> None:
    (tmp_path / "pubify.yaml").write_text(
        "pubify-ppt:\n"
        "  presentations_root: decks\n",
        encoding="utf-8",
    )

    init_workspace(tmp_path)

    assert (tmp_path / "pubify.yaml").read_text(encoding="utf-8") == (
        "pubify-ppt:\n"
        "  presentations_root: decks\n"
    )
    assert (tmp_path / "decks").is_dir()


def test_init_presentation_creates_starter_files_and_deck_anchors(tmp_path: Path) -> None:
    init_workspace(tmp_path)

    presentation_root = init_presentation_by_id(tmp_path, "demo")

    assert presentation_root == tmp_path / "slides" / "demo"
    assert (presentation_root / "ppt.yaml").read_text(encoding="utf-8").startswith("deck: deck.pptx\n")
    figures_source = (presentation_root / "figures.py").read_text(encoding="utf-8")
    assert '@data("example.csv")' in figures_source
    assert "def plot_example" in figures_source
    assert "def compute_example" in figures_source
    assert (presentation_root / "data" / "example.csv").read_text(encoding="utf-8") == "x,y\n1,1\n2,4\n3,9\n"

    deck = Presentation(presentation_root / "deck.pptx")
    slide = deck.slides[0]
    descriptions = [_shape_alt_text(shape) for shape in slide.shapes]
    text = "\n".join(shape.text for shape in slide.shapes if hasattr(shape, "text"))
    assert "{{fig:example}}" in descriptions
    assert "{{stat:example.count}}" in text


def test_init_presentation_preserves_existing_user_files(tmp_path: Path) -> None:
    init_workspace(tmp_path)
    presentation_root = tmp_path / "slides" / "demo"
    presentation_root.mkdir(parents=True)
    (presentation_root / "ppt.yaml").write_text("deck: custom.pptx\n", encoding="utf-8")
    (presentation_root / "figures.py").write_text("# custom\n", encoding="utf-8")
    custom_deck = Presentation()
    custom_deck.slides.add_slide(custom_deck.slide_layouts[6])
    custom_deck.save(presentation_root / "deck.pptx")

    init_presentation_by_id(tmp_path, "demo")

    assert (presentation_root / "ppt.yaml").read_text(encoding="utf-8") == "deck: custom.pptx\n"
    assert (presentation_root / "figures.py").read_text(encoding="utf-8") == "# custom\n"
    assert len(Presentation(presentation_root / "deck.pptx").slides) == 1


def test_init_presentation_preserves_symlinked_data_root(tmp_path: Path) -> None:
    init_workspace(tmp_path)
    external_data = tmp_path / "external-data"
    external_data.mkdir()
    presentation_root = tmp_path / "slides" / "demo"
    presentation_root.mkdir(parents=True)
    (presentation_root / "data").symlink_to(external_data, target_is_directory=True)

    init_presentation_by_id(tmp_path, "demo")

    assert (presentation_root / "data").is_symlink()
    assert (external_data / "ppt-artifacts" / "figures").is_dir()
    assert (external_data / "ppt-artifacts" / "backups").is_dir()


def test_init_presentation_rejects_path_like_ids(tmp_path: Path) -> None:
    init_workspace(tmp_path)

    with pytest.raises(ValueError, match="single non-empty path segment"):
        init_presentation_by_id(tmp_path, ".")
    with pytest.raises(ValueError, match="single non-empty path segment"):
        init_presentation_by_id(tmp_path, "../demo")


def _shape_alt_text(shape: object) -> str | None:
    c_nv_pr = shape._element.xpath(".//p:cNvPr")[0]
    return c_nv_pr.get("descr")
