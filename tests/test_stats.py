from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.util import Inches
import pytest

from pubify_ppt.discovery import load_presentation_definition
from pubify_ppt.init import init_presentation_by_id, init_workspace
from pubify_ppt.stats import update_stats
from pubify_ppt.anchors import shape_alt_text


def test_update_stats_replaces_starter_dictionary_token(tmp_path: Path) -> None:
    init_workspace(tmp_path)
    init_presentation_by_id(tmp_path, "demo")
    presentation = load_presentation_definition(tmp_path, "demo")

    replacements = update_stats(presentation)

    assert [(item.token, item.value) for item in replacements] == [("{{stat:example.count}}", "3")]
    assert "Example count: 3" in _deck_text(tmp_path / "slides" / "demo" / "deck.pptx")
    assert "{{stat:example.count=3}}" in _shape_alt_texts(tmp_path / "slides" / "demo" / "deck.pptx")


def test_update_stats_replaces_local_wrapper_around_source_publication_stat(tmp_path: Path) -> None:
    init_workspace(tmp_path)
    init_presentation_by_id(tmp_path, "demo")
    source_root = tmp_path / "papers" / "ao4elt8"
    source_data = source_root / "data"
    source_data.mkdir(parents=True)
    (source_data / "value.txt").write_text("source", encoding="utf-8")
    (source_root / "figures.py").write_text(
        "\n".join(
            [
                "from pubify_data import data, stat",
                "@data('value.txt')",
                "def load_value(ctx, path):",
                "    return path.read_text(encoding='utf-8')",
                "@stat",
                "def compute_summary(ctx, value):",
                "    return value",
            ]
        )
        + "\n",
        encoding="utf-8",
    )
    presentation_root = tmp_path / "slides" / "demo"
    (presentation_root / "figures.py").write_text(
        "\n".join(
            [
                "from pubify_data import stat",
                "from pubify_ppt import StatResult",
                "@stat",
                "def compute_source_summary(ctx):",
                "    return StatResult(ctx.source('ao4elt8').stat('summary').values[0].value)",
            ]
        )
        + "\n",
        encoding="utf-8",
    )
    (presentation_root / "ppt.yaml").write_text(
        "\n".join(
            [
                "deck: deck.pptx",
                "backup_retention: 5",
                "defaults:",
                "  image_format: png",
                "  dpi: 200",
                "  fit: contain",
                "external_data_roots:",
                "sources:",
                "  ao4elt8: papers/ao4elt8",
            ]
        )
        + "\n",
        encoding="utf-8",
    )
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(6.0), Inches(0.5)).text = "Value: {{stat:source_summary}}"
    deck.save(presentation_root / "deck.pptx")
    presentation = load_presentation_definition(tmp_path, "demo")

    replacements = update_stats(presentation, stat_id="source_summary")

    assert [(item.token, item.value) for item in replacements] == [("{{stat:source_summary}}", "source")]
    assert "Value: source" in _deck_text(presentation_root / "deck.pptx")
    assert "{{stat:source_summary=source}}" in _shape_alt_texts(presentation_root / "deck.pptx")


def test_update_stats_can_refresh_previously_replaced_token(tmp_path: Path) -> None:
    init_workspace(tmp_path)
    init_presentation_by_id(tmp_path, "demo")
    presentation_root = tmp_path / "slides" / "demo"
    figures_path = presentation_root / "figures.py"
    figures_path.write_text(_stat_module("first"), encoding="utf-8")
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(6.0), Inches(0.5)).text = "Total: {{stat:total}}"
    deck.save(presentation_root / "deck.pptx")
    presentation = load_presentation_definition(tmp_path, "demo")

    first_replacements = update_stats(presentation)

    assert [(item.token, item.value) for item in first_replacements] == [("{{stat:total}}", "first")]
    assert "Total: first" in _deck_text(presentation_root / "deck.pptx")

    figures_path.write_text(_stat_module("second"), encoding="utf-8")
    presentation = load_presentation_definition(tmp_path, "demo")

    second_replacements = update_stats(presentation)

    assert [(item.token, item.value) for item in second_replacements] == [("{{stat:total}}", "second")]
    assert "Total: second" in _deck_text(presentation_root / "deck.pptx")


def test_update_stats_can_refresh_split_run_token(tmp_path: Path) -> None:
    init_workspace(tmp_path)
    init_presentation_by_id(tmp_path, "demo")
    presentation_root = tmp_path / "slides" / "demo"
    figures_path = presentation_root / "figures.py"
    figures_path.write_text(_stat_module("first"), encoding="utf-8")
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    text_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(6.0), Inches(0.5))
    paragraph = text_box.text_frame.paragraphs[0]
    paragraph.add_run().text = "Total: "
    paragraph.add_run().text = "{{"
    paragraph.add_run().text = "stat:total"
    paragraph.add_run().text = "}}"
    deck.save(presentation_root / "deck.pptx")
    presentation = load_presentation_definition(tmp_path, "demo")

    first_replacements = update_stats(presentation)

    assert [(item.token, item.value) for item in first_replacements] == [("{{stat:total}}", "first")]
    deck = Presentation(presentation_root / "deck.pptx")
    paragraph = deck.slides[0].shapes[0].text_frame.paragraphs[0]
    assert [run.text for run in paragraph.runs] == ["Total: ", "first", "", ""]

    figures_path.write_text(_stat_module("second"), encoding="utf-8")
    presentation = load_presentation_definition(tmp_path, "demo")

    second_replacements = update_stats(presentation)

    assert [(item.token, item.value) for item in second_replacements] == [("{{stat:total}}", "second")]
    assert "Total: second" in _deck_text(presentation_root / "deck.pptx")


def test_update_stats_replaces_scalar_token_and_preserves_run_format(tmp_path: Path) -> None:
    init_workspace(tmp_path)
    init_presentation_by_id(tmp_path, "demo")
    presentation_root = tmp_path / "slides" / "demo"
    (presentation_root / "figures.py").write_text(
        "\n".join(
            [
                "from pubify_data import stat",
                "@stat",
                "def compute_total(ctx):",
                "    return 42",
            ]
        )
        + "\n",
        encoding="utf-8",
    )
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    text_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(5.0), Inches(0.5))
    run = text_box.text_frame.paragraphs[0].add_run()
    run.text = "{{stat:total}}"
    run.font.bold = True
    deck.save(presentation_root / "deck.pptx")
    presentation = load_presentation_definition(tmp_path, "demo")

    replacements = update_stats(presentation)

    assert [item.value for item in replacements] == ["42"]
    deck = Presentation(presentation_root / "deck.pptx")
    run = deck.slides[0].shapes[0].text_frame.paragraphs[0].runs[0]
    assert run.text == "42"
    assert run.font.bold is True
    assert shape_alt_text(deck.slides[0].shapes[0]) == "{{stat:total=42}}"


def test_update_stats_targeted_update_leaves_other_tokens_unchanged(tmp_path: Path) -> None:
    init_workspace(tmp_path)
    init_presentation_by_id(tmp_path, "demo")
    presentation_root = tmp_path / "slides" / "demo"
    (presentation_root / "figures.py").write_text(
        "\n".join(
            [
                "from pubify_data import stat",
                "@stat",
                "def compute_first(ctx):",
                "    return 'one'",
                "@stat",
                "def compute_second(ctx):",
                "    return 'two'",
            ]
        )
        + "\n",
        encoding="utf-8",
    )
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(6.0), Inches(0.5)).text = "{{stat:first}}"
    slide.shapes.add_textbox(Inches(0.5), Inches(1.0), Inches(6.0), Inches(0.5)).text = "{{stat:second}}"
    deck.save(presentation_root / "deck.pptx")
    presentation = load_presentation_definition(tmp_path, "demo")

    update_stats(presentation, stat_id="first")

    text = _deck_text(presentation_root / "deck.pptx")
    assert "one" in text
    assert "{{stat:second}}" in text


def test_update_stats_targeted_refresh_preserves_previous_rendered_values(tmp_path: Path) -> None:
    init_workspace(tmp_path)
    init_presentation_by_id(tmp_path, "demo")
    presentation_root = tmp_path / "slides" / "demo"
    (presentation_root / "figures.py").write_text(
        "\n".join(
            [
                "from pubify_data import stat",
                "@stat",
                "def compute_first(ctx):",
                "    return 'one'",
                "@stat",
                "def compute_second(ctx):",
                "    return 'two'",
            ]
        )
        + "\n",
        encoding="utf-8",
    )
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(6.0), Inches(0.5)).text = "{{stat:first}}"
    slide.shapes.add_textbox(Inches(0.5), Inches(1.0), Inches(6.0), Inches(0.5)).text = "{{stat:second}}"
    deck.save(presentation_root / "deck.pptx")
    presentation = load_presentation_definition(tmp_path, "demo")

    update_stats(presentation, stat_id="first")
    update_stats(presentation, stat_id="second")

    text = _deck_text(presentation_root / "deck.pptx")
    assert "one" in text
    assert "two" in text


def test_update_stats_reports_missing_previous_rendered_value(tmp_path: Path) -> None:
    init_workspace(tmp_path)
    init_presentation_by_id(tmp_path, "demo")
    presentation_root = tmp_path / "slides" / "demo"
    figures_path = presentation_root / "figures.py"
    figures_path.write_text(_stat_module("first"), encoding="utf-8")
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(6.0), Inches(0.5)).text = "Total: {{stat:total}}"
    deck.save(presentation_root / "deck.pptx")
    presentation = load_presentation_definition(tmp_path, "demo")
    update_stats(presentation)

    deck = Presentation(presentation_root / "deck.pptx")
    deck.slides[0].shapes[0].text = "Total: edited"
    deck.save(presentation_root / "deck.pptx")
    figures_path.write_text(_stat_module("second"), encoding="utf-8")
    presentation = load_presentation_definition(tmp_path, "demo")

    with pytest.raises(ValueError, match="Restore it or reinsert"):
        update_stats(presentation)


def test_update_stats_can_repair_by_reinserting_token(tmp_path: Path) -> None:
    init_workspace(tmp_path)
    init_presentation_by_id(tmp_path, "demo")
    presentation_root = tmp_path / "slides" / "demo"
    figures_path = presentation_root / "figures.py"
    figures_path.write_text(_stat_module("first"), encoding="utf-8")
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(6.0), Inches(0.5)).text = "Total: {{stat:total}}"
    deck.save(presentation_root / "deck.pptx")
    presentation = load_presentation_definition(tmp_path, "demo")
    update_stats(presentation)

    deck = Presentation(presentation_root / "deck.pptx")
    deck.slides[0].shapes[0].text = "Total: {{stat:total}}"
    deck.save(presentation_root / "deck.pptx")
    figures_path.write_text(_stat_module("second"), encoding="utf-8")
    presentation = load_presentation_definition(tmp_path, "demo")

    update_stats(presentation)

    assert "Total: second" in _deck_text(presentation_root / "deck.pptx")


def test_update_stats_reports_missing_dictionary_key(tmp_path: Path) -> None:
    init_workspace(tmp_path)
    init_presentation_by_id(tmp_path, "demo")
    presentation_root = tmp_path / "slides" / "demo"
    deck = Presentation(presentation_root / "deck.pptx")
    slide = deck.slides[0]
    slide.shapes.add_textbox(Inches(0.5), Inches(5.5), Inches(4.0), Inches(0.5)).text = "{{stat:example.missing}}"
    deck.save(presentation_root / "deck.pptx")
    presentation = load_presentation_definition(tmp_path, "demo")

    with pytest.raises(ValueError, match="Missing key 'missing' for stat 'example'"):
        update_stats(presentation)


def _deck_text(deck_path: Path) -> str:
    deck = Presentation(deck_path)
    return "\n".join(shape.text for slide in deck.slides for shape in slide.shapes if hasattr(shape, "text"))


def _shape_alt_texts(deck_path: Path) -> list[str]:
    deck = Presentation(deck_path)
    return [
        value
        for slide in deck.slides
        for shape in slide.shapes
        if (value := shape_alt_text(shape)) is not None
    ]


def _stat_module(value: str) -> str:
    return "\n".join(
        [
            "from pubify_data import stat",
            "@stat",
            "def compute_total(ctx):",
            f"    return {value!r}",
        ]
    ) + "\n"
