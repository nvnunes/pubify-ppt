from __future__ import annotations

from pathlib import Path
from xml.etree import ElementTree
from zipfile import ZipFile

import matplotlib.pyplot as plt
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches

from pubify_ppt.discovery import load_presentation_definition
from pubify_ppt.init import init_presentation_by_id, init_workspace
from pubify_ppt.update import update_presentation


def test_full_update_patches_only_touched_slide_rels_and_media_parts(tmp_path: Path) -> None:
    init_workspace(tmp_path)
    init_presentation_by_id(tmp_path, "demo")
    presentation_root = tmp_path / "slides" / "demo"
    deck_path = presentation_root / "deck.pptx"
    deck = Presentation(deck_path)
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(2), Inches(0.5)).text = "untouched"
    deck.save(deck_path)
    before = _package_payloads(deck_path)
    presentation = load_presentation_definition(tmp_path, "demo")
    plt.close("all")

    update_presentation(presentation)

    changed = _changed_package_entries(before, _package_payloads(deck_path))
    assert changed <= {
        "[Content_Types].xml",
        "ppt/slides/slide1.xml",
        "ppt/slides/_rels/slide1.xml.rels",
        "ppt/media/image1.png",
    }
    assert "ppt/slides/slide2.xml" not in changed


def test_full_update_refreshes_existing_picture_anchor_as_media_only(tmp_path: Path) -> None:
    init_workspace(tmp_path)
    init_presentation_by_id(tmp_path, "demo")
    presentation_root = tmp_path / "slides" / "demo"
    deck_path = presentation_root / "deck.pptx"
    _write_line_figure_module(presentation_root, [1, 2, 3])
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.5), Inches(0.5), Inches(3), Inches(2))
    _set_shape_alt_text(shape, "{{fig:line}}")
    deck.save(deck_path)
    presentation = load_presentation_definition(tmp_path, "demo")
    update_presentation(presentation)
    before_second_update = _package_payloads(deck_path)

    _write_line_figure_module(presentation_root, [3, 1, 2])
    presentation = load_presentation_definition(tmp_path, "demo")
    update_presentation(presentation)

    changed = _changed_package_entries(before_second_update, _package_payloads(deck_path))
    assert changed == {"ppt/media/image1.png"}


def test_full_update_detaches_shared_managed_picture_relationships(tmp_path: Path) -> None:
    init_workspace(tmp_path)
    init_presentation_by_id(tmp_path, "demo")
    presentation_root = tmp_path / "slides" / "demo"
    deck_path = presentation_root / "deck.pptx"
    _write_two_line_figure_module(presentation_root)
    seed = presentation_root / "seed.png"

    Image.new("RGB", (20, 20), "blue").save(seed)
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    first = slide.shapes.add_picture(str(seed), Inches(0.5), Inches(0.5), width=Inches(2))
    second = slide.shapes.add_picture(str(seed), Inches(3), Inches(0.5), width=Inches(2))
    _set_shape_alt_text(first, "{{fig:first}}")
    _set_shape_alt_text(second, "{{fig:second}}")
    deck.save(deck_path)
    presentation = load_presentation_definition(tmp_path, "demo")

    update_presentation(presentation)

    assert len(set(_picture_blip_ids(deck_path, slide_number=1))) == 2


def _package_payloads(deck_path: Path) -> dict[str, bytes]:
    with ZipFile(deck_path) as package:
        return {name: package.read(name) for name in package.namelist()}


def _changed_package_entries(before: dict[str, bytes], after: dict[str, bytes]) -> set[str]:
    names = set(before) | set(after)
    return {name for name in names if before.get(name) != after.get(name)}


def _write_line_figure_module(presentation_root: Path, y_values: list[int]) -> None:
    (presentation_root / "figures.py").write_text(
        "\n".join(
            [
                "import matplotlib.pyplot as plt",
                "from pubify_data import figure",
                "@figure",
                "def plot_line(ctx):",
                "    fig, ax = plt.subplots()",
                f"    ax.plot([1, 2, 3], {y_values!r})",
                "    return fig",
            ]
        )
        + "\n",
        encoding="utf-8",
    )


def _write_two_line_figure_module(presentation_root: Path) -> None:
    (presentation_root / "figures.py").write_text(
        "\n".join(
            [
                "import matplotlib.pyplot as plt",
                "from pubify_data import figure",
                "@figure",
                "def plot_first(ctx):",
                "    fig, ax = plt.subplots()",
                "    ax.plot([1, 2, 3], [1, 2, 3])",
                "    return fig",
                "@figure",
                "def plot_second(ctx):",
                "    fig, ax = plt.subplots()",
                "    ax.plot([1, 2, 3], [3, 1, 2])",
                "    return fig",
            ]
        )
        + "\n",
        encoding="utf-8",
    )


def _set_shape_alt_text(shape: object, value: str) -> None:
    c_nv_pr = shape._element.xpath(".//p:cNvPr")[0]
    c_nv_pr.set("descr", value)


def _picture_blip_ids(deck_path: Path, *, slide_number: int) -> list[str]:
    relationship_attribute_namespace = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
    with ZipFile(deck_path) as package:
        slide_xml = package.read(f"ppt/slides/slide{slide_number}.xml")
    slide = ElementTree.fromstring(slide_xml)
    return [
        value
        for element in slide.iter()
        if element.tag.endswith("}blip")
        if (value := element.attrib.get(f"{relationship_attribute_namespace}embed")) is not None
    ]
