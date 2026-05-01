"""PowerPoint anchor and inline-token discovery for ``pubify-ppt``."""

from __future__ import annotations

from dataclasses import dataclass
import re

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn
from pptx.presentation import Presentation as PresentationObject


FIGURE_TOKEN_RE = re.compile(r"^\{\{fig:([A-Za-z0-9_.-]+)(?::([1-9][0-9]*))?\}\}$")
STAT_TOKEN_RE = re.compile(r"\{\{stat:([A-Za-z0-9_.-]+)\}\}")
STAT_RENDERED_TOKEN_RE = re.compile(r"^\{\{stat:([A-Za-z0-9_.-]+)=(.*?)\}\}$", re.DOTALL)
TABLE_TOKEN_RE = re.compile(r"^\{\{table:([A-Za-z0-9_.-]+)\}\}$")
MANAGED_TOKEN_RE = re.compile(r"\{\{([^{}]+)\}\}")


@dataclass(frozen=True)
class FigureAnchor:
    """One PowerPoint shape carrying a managed figure token."""

    slide_number: int
    shape_index: int
    figure_id: str
    panel_number: int | None
    token: str
    shape: object


@dataclass(frozen=True)
class StatToken:
    """One inline stat token discovered in PowerPoint text."""

    slide_number: int
    stat_id: str
    key: str | None
    token: str


@dataclass(frozen=True)
class TableAnchor:
    """One PowerPoint shape carrying a managed table token."""

    slide_number: int
    shape_index: int
    table_id: str
    token: str
    shape: object


def discover_figure_anchors(deck_path: object) -> tuple[FigureAnchor, ...]:
    """Discover figure anchors from PowerPoint shape alt text."""

    deck = Presentation(deck_path)
    return discover_figure_anchors_in_deck(deck)


def discover_figure_anchors_in_deck(deck: PresentationObject) -> tuple[FigureAnchor, ...]:
    """Discover figure anchors from an already-open PowerPoint deck."""

    anchors: list[FigureAnchor] = []
    for slide_number, slide in enumerate(deck.slides, start=1):
        for shape_index, shape in enumerate(slide.shapes, start=1):
            token = _figure_anchor_token(shape)
            if token is None:
                continue
            match = FIGURE_TOKEN_RE.fullmatch(token)
            if match is None:
                continue
            panel = int(match.group(2)) if match.group(2) is not None else None
            anchors.append(FigureAnchor(slide_number, shape_index, match.group(1), panel, match.group(0), shape))
    return tuple(anchors)


def discover_stat_tokens(deck_path: object) -> tuple[StatToken, ...]:
    """Discover supported stat tokens from PowerPoint text shapes."""

    deck = Presentation(deck_path)
    tokens: list[StatToken] = []
    for slide_number, slide in enumerate(deck.slides, start=1):
        for shape in slide.shapes:
            text_frame = getattr(shape, "text_frame", None)
            if text_frame is None:
                continue
            for paragraph in text_frame.paragraphs:
                for run in paragraph.runs:
                    for match in STAT_TOKEN_RE.finditer(run.text):
                        stat_id, key = split_stat_reference(match.group(1))
                        tokens.append(StatToken(slide_number, stat_id, key, match.group(0)))
    return tuple(tokens)


def discover_table_anchors_in_deck(deck: PresentationObject) -> tuple[TableAnchor, ...]:
    """Discover table anchors from an already-open PowerPoint deck."""

    anchors: list[TableAnchor] = []
    for slide_number, slide in enumerate(deck.slides, start=1):
        for shape_index, shape in enumerate(slide.shapes, start=1):
            token = _table_anchor_token(shape)
            if token is None:
                continue
            match = TABLE_TOKEN_RE.fullmatch(token)
            if match is None:
                continue
            anchors.append(TableAnchor(slide_number, shape_index, match.group(1), match.group(0), shape))
    return tuple(anchors)


def validate_deck_anchors(
    deck_path: object,
    *,
    figure_ids: set[str],
    stat_ids: set[str],
    table_ids: set[str],
    report_shared_figure_relationships: bool = True,
) -> list[str]:
    """Return static deck-anchor validation errors without mutating the deck."""

    deck = Presentation(deck_path)
    errors: list[str] = []
    for slide_number, slide in enumerate(deck.slides, start=1):
        managed_picture_relationships: dict[str, list[str]] = {}
        for shape_index, shape in enumerate(slide.shapes, start=1):
            alt_text = shape_alt_text(shape)
            alt_figure_token: str | None = None
            alt_table_token: str | None = None
            if alt_text:
                stripped_alt_text = alt_text.strip()
                if FIGURE_TOKEN_RE.fullmatch(stripped_alt_text) is not None:
                    alt_figure_token = stripped_alt_text
                    _record_managed_picture_relationship(
                        managed_picture_relationships,
                        shape,
                        token=stripped_alt_text,
                    )
                if TABLE_TOKEN_RE.fullmatch(stripped_alt_text) is not None:
                    alt_table_token = stripped_alt_text
                _validate_alt_text(
                    errors,
                    shape,
                    slide_number=slide_number,
                    alt_text=stripped_alt_text,
                    figure_ids=figure_ids,
                    stat_ids=stat_ids,
                    table_ids=table_ids,
                )
            _validate_text_shape(
                errors,
                shape,
                slide_number=slide_number,
                figure_ids=figure_ids,
                stat_ids=stat_ids,
                table_ids=table_ids,
                alt_figure_token=alt_figure_token,
                alt_table_token=alt_table_token,
            )
        if report_shared_figure_relationships:
            _validate_shared_managed_picture_relationships(
                errors,
                slide_number=slide_number,
                relationships=managed_picture_relationships,
            )
    return errors


def shape_alt_text(shape: object) -> str | None:
    """Return a shape's description alt text as exposed in the Open XML tree."""

    matches = shape._element.xpath(".//p:cNvPr")
    if not matches:
        return None
    return matches[0].get("descr")


def set_shape_alt_text(shape: object, value: str) -> None:
    """Set a shape's description alt text in the Open XML tree."""

    matches = shape._element.xpath(".//p:cNvPr")
    if not matches:
        raise ValueError("Shape does not expose a PowerPoint non-visual properties node")
    matches[0].set("descr", value)


def _record_managed_picture_relationship(
    relationships: dict[str, list[str]],
    shape: object,
    *,
    token: str,
) -> None:
    r_ids = _embedded_image_r_ids(shape._element)
    if len(r_ids) != 1:
        return
    relationships.setdefault(r_ids[0], []).append(token)


def _validate_shared_managed_picture_relationships(
    errors: list[str],
    *,
    slide_number: int,
    relationships: dict[str, list[str]],
) -> None:
    for r_id, tokens in sorted(relationships.items()):
        unique_tokens = sorted(set(tokens))
        if len(unique_tokens) <= 1:
            continue
        joined = ", ".join(unique_tokens)
        errors.append(
            f"Slide {slide_number}: managed figure picture anchors share image relationship {r_id}: {joined}; "
            "rerun update to detach copied placeholders"
        )


def _embedded_image_r_ids(element: object) -> tuple[str, ...]:
    return tuple(
        r_id
        for blip in element.xpath(".//a:blip")
        if (r_id := blip.get(qn("r:embed"))) is not None
    )


def _validate_alt_text(
    errors: list[str],
    shape: object,
    *,
    slide_number: int,
    alt_text: str,
    figure_ids: set[str],
    stat_ids: set[str],
    table_ids: set[str],
) -> None:
    managed_match = MANAGED_TOKEN_RE.fullmatch(alt_text)
    if managed_match is None:
        return
    if managed_match.group(1).startswith("stat:"):
        _validate_stat_alt_text(errors, slide_number=slide_number, alt_text=alt_text, stat_ids=stat_ids)
        return
    if managed_match.group(1).startswith("table:"):
        _validate_table_anchor_token(
            errors,
            shape,
            slide_number=slide_number,
            token=alt_text,
            table_ids=table_ids,
        )
        return
    figure_match = FIGURE_TOKEN_RE.fullmatch(alt_text)
    if figure_match is None:
        if managed_match.group(1).startswith("fig:"):
            errors.append(f"Slide {slide_number}: malformed figure anchor token {alt_text!r}")
        return

    _validate_figure_anchor_token(
        errors,
        shape,
        slide_number=slide_number,
        token=figure_match.group(0),
        figure_ids=figure_ids,
    )


def _validate_stat_alt_text(
    errors: list[str],
    *,
    slide_number: int,
    alt_text: str,
    stat_ids: set[str],
) -> None:
    rendered_match = STAT_RENDERED_TOKEN_RE.fullmatch(alt_text)
    token_match = STAT_TOKEN_RE.fullmatch(alt_text)
    match = rendered_match if rendered_match is not None else token_match
    if match is None:
        errors.append(f"Slide {slide_number}: malformed stat anchor token {alt_text!r}")
        return
    stat_id, _ = split_stat_reference(match.group(1), stat_ids=stat_ids)
    if stat_id not in stat_ids:
        errors.append(f"Slide {slide_number}: unknown stat anchor {alt_text}")


def _validate_text_shape(
    errors: list[str],
    shape: object,
    *,
    slide_number: int,
    figure_ids: set[str],
    stat_ids: set[str],
    table_ids: set[str],
    alt_figure_token: str | None,
    alt_table_token: str | None,
) -> None:
    text_frame = getattr(shape, "text_frame", None)
    if text_frame is None:
        return
    visible_figure_token = _visible_figure_anchor_token(shape)
    visible_table_token = _visible_table_anchor_token(shape)
    if visible_figure_token is not None:
        if alt_figure_token is None:
            _validate_figure_anchor_token(
                errors,
                shape,
                slide_number=slide_number,
                token=visible_figure_token,
                figure_ids=figure_ids,
            )
        elif alt_figure_token != visible_figure_token:
            errors.append(
                "Slide "
                f"{slide_number}: figure anchor alt text {alt_figure_token!r} conflicts with "
                f"visible figure token {visible_figure_token!r}"
            )
    elif visible_table_token is not None:
        if alt_table_token is None:
            _validate_table_anchor_token(
                errors,
                shape,
                slide_number=slide_number,
                token=visible_table_token,
                table_ids=table_ids,
            )
        elif alt_table_token != visible_table_token:
            errors.append(
                "Slide "
                f"{slide_number}: table anchor alt text {alt_table_token!r} conflicts with "
                f"visible table token {visible_table_token!r}"
            )
    else:
        text = _shape_text(shape).strip()
        managed_match = MANAGED_TOKEN_RE.fullmatch(text)
        if managed_match is not None and managed_match.group(1).startswith("fig:"):
            errors.append(f"Slide {slide_number}: malformed figure anchor token {text!r}")
        if managed_match is not None and managed_match.group(1).startswith("table:"):
            errors.append(f"Slide {slide_number}: malformed table anchor token {text!r}")
    for paragraph in text_frame.paragraphs:
        paragraph_text = "".join(run.text for run in paragraph.runs)
        stat_tokens = [match.group(0) for match in STAT_TOKEN_RE.finditer(paragraph_text)]
        if len(stat_tokens) > 1:
            errors.append(
                f"Slide {slide_number}: stat-managed text boxes must contain exactly one stat token; "
                "split multiple stats into separate text boxes"
            )
        for managed_match in MANAGED_TOKEN_RE.finditer(paragraph_text):
            token = managed_match.group(0)
            if not managed_match.group(1).startswith("stat:"):
                continue
            stat_match = STAT_TOKEN_RE.fullmatch(token)
            if stat_match is None:
                errors.append(f"Slide {slide_number}: malformed stat token {token!r}")
                continue
            stat_id, _ = split_stat_reference(stat_match.group(1), stat_ids=stat_ids)
            if stat_id not in stat_ids:
                errors.append(f"Slide {slide_number}: unknown stat token {token}")


def _validate_figure_anchor_token(
    errors: list[str],
    shape: object,
    *,
    slide_number: int,
    token: str,
    figure_ids: set[str],
) -> None:
    match = FIGURE_TOKEN_RE.fullmatch(token)
    if match is None:
        errors.append(f"Slide {slide_number}: malformed figure anchor token {token!r}")
        return
    figure_id = match.group(1)
    if figure_id not in figure_ids:
        errors.append(f"Slide {slide_number}: unknown figure anchor {token}")

    unsupported = _unsupported_figure_anchor_features(shape)
    if unsupported:
        errors.append(f"Slide {slide_number}: unsupported figure anchor {token}: {unsupported}")


def _validate_table_anchor_token(
    errors: list[str],
    shape: object,
    *,
    slide_number: int,
    token: str,
    table_ids: set[str],
) -> None:
    match = TABLE_TOKEN_RE.fullmatch(token)
    if match is None:
        errors.append(f"Slide {slide_number}: malformed table anchor token {token!r}")
        return
    table_id = match.group(1)
    if table_id not in table_ids:
        errors.append(f"Slide {slide_number}: unknown table anchor {token}")

    unsupported = _unsupported_table_anchor_features(shape)
    if unsupported:
        errors.append(f"Slide {slide_number}: unsupported table anchor {token}: {unsupported}")


def _unsupported_figure_anchor_features(shape: object) -> str | None:
    features: list[str] = []
    if getattr(shape, "shape_type", None) == MSO_SHAPE_TYPE.GROUP:
        features.append("grouped shape")
    rotation = getattr(shape, "rotation", 0)
    if rotation not in (0, 0.0):
        features.append("rotation")
    if getattr(shape, "shape_type", None) == MSO_SHAPE_TYPE.PICTURE:
        crop_values = (
            getattr(shape, "crop_left", 0),
            getattr(shape, "crop_right", 0),
            getattr(shape, "crop_top", 0),
            getattr(shape, "crop_bottom", 0),
        )
        if any(value not in (0, 0.0) for value in crop_values):
            features.append("crop")
    return ", ".join(features) if features else None


def _unsupported_table_anchor_features(shape: object) -> str | None:
    features: list[str] = []
    if getattr(shape, "shape_type", None) == MSO_SHAPE_TYPE.GROUP:
        features.append("grouped shape")
    if getattr(shape, "has_chart", False):
        features.append("chart")
    if not getattr(shape, "has_table", False) and getattr(shape, "text_frame", None) is None:
        features.append("unsupported shape type")
    rotation = getattr(shape, "rotation", 0)
    if rotation not in (0, 0.0):
        features.append("rotation")
    return ", ".join(features) if features else None


def _figure_anchor_token(shape: object) -> str | None:
    alt_text = shape_alt_text(shape)
    if alt_text is not None and FIGURE_TOKEN_RE.fullmatch(alt_text.strip()) is not None:
        return alt_text.strip()
    return _visible_figure_anchor_token(shape)


def _visible_figure_anchor_token(shape: object) -> str | None:
    text = _shape_text(shape).strip()
    if not text:
        return None
    if FIGURE_TOKEN_RE.fullmatch(text) is None:
        return None
    return text


def _table_anchor_token(shape: object) -> str | None:
    alt_text = shape_alt_text(shape)
    if alt_text is not None and TABLE_TOKEN_RE.fullmatch(alt_text.strip()) is not None:
        return alt_text.strip()
    return _visible_table_anchor_token(shape)


def _visible_table_anchor_token(shape: object) -> str | None:
    text = _shape_text(shape).strip()
    if not text:
        return None
    if TABLE_TOKEN_RE.fullmatch(text) is None:
        return None
    return text


def _shape_text(shape: object) -> str:
    text_frame = getattr(shape, "text_frame", None)
    if text_frame is None:
        return ""
    return "\n".join("".join(run.text for run in paragraph.runs) for paragraph in text_frame.paragraphs)


def split_stat_reference(reference: str, *, stat_ids: set[str] | tuple[str, ...] | None = None) -> tuple[str, str | None]:
    """Split a stat token body into ``(stat_id, key)`` using known ids when available."""

    if stat_ids:
        for stat_id in sorted(stat_ids, key=len, reverse=True):
            if reference == stat_id:
                return stat_id, None
            prefix = f"{stat_id}."
            if reference.startswith(prefix):
                key = reference[len(prefix) :]
                return stat_id, key or None
    if "." in reference:
        stat_id, key = reference.split(".", 1)
        return stat_id, key
    return reference, None
