from __future__ import annotations

from collections.abc import Mapping
from dataclasses import dataclass
import math
from numbers import Real
from pathlib import Path
import re

from matplotlib.transforms import Bbox
from PIL import Image
from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.presentation import Presentation as PresentationObject
from pptx.util import Emu
import pubify_data
import pubify_mpl

from pubify_ppt.anchors import (
    FigureAnchor,
    discover_figure_anchors_in_deck,
    set_shape_alt_text,
)
from pubify_ppt.backups import write_patched_deck
from pubify_ppt.discovery import PresentationDefinition
from pubify_ppt.ooxml import MediaReplacement
from pubify_ppt.runtime import check_presentation, ensure_generated_artifact_paths


EMU_PER_INCH = 914400
FIGURE_PREPARATION_OPTIONS = {
    "style",
    "keep_titles",
    "hide_labels",
    "hide_annotations",
    "hide_ticks",
    "hide_tick_labels",
    "hide_grid",
    "hide_cbar",
    "skip_clone",
    "extra_rcparams",
    "prepare_export",
    "export_pad_inches",
    "export_pad_left_inches",
    "export_pad_right_inches",
    "export_pad_top_inches",
    "export_pad_bottom_inches",
}
EXPORT_PADDING_OPTIONS = {
    "export_pad_inches",
    "export_pad_left_inches",
    "export_pad_right_inches",
    "export_pad_top_inches",
    "export_pad_bottom_inches",
}
ASPECT_RATIO_TOLERANCE = 0.02


@dataclass(frozen=True)
class FigureOutput:
    """One rendered figure output and the slide it was applied to."""

    slide_number: int
    shape_index: int
    token: str
    path: Path


@dataclass(frozen=True)
class FigureUpdateResult:
    """Figure update changes applied to an open deck."""

    deck: PresentationObject
    outputs: tuple[FigureOutput, ...]
    touched_slide_numbers: tuple[int, ...]
    media_replacements: tuple[MediaReplacement, ...]


@dataclass(frozen=True)
class FigureAnchorUpdate:
    """One figure anchor update and the package parts it requires."""

    media_replacement: MediaReplacement | None
    touches_slide_xml: bool


@dataclass(frozen=True)
class _ExportPadding:
    """PowerPoint figure export padding in inches."""

    symmetric: float
    left: float
    right: float
    top: float
    bottom: float
    has_side_specific: bool


def update_figures(
    presentation: PresentationDefinition,
    *,
    figure_id: str | None = None,
) -> tuple[Path, ...]:
    """Render selected figures and replace their PowerPoint anchors in place."""

    result = update_figures_in_deck(presentation, figure_id=figure_id)
    write_patched_deck(
        presentation,
        result.deck,
        touched_slide_numbers=result.touched_slide_numbers,
        media_replacements=result.media_replacements,
    )
    return tuple(output.path for output in result.outputs)


def update_figures_to_output(
    presentation: PresentationDefinition,
    *,
    figure_id: str | None = None,
    output: Path | None = None,
) -> tuple[Path, ...]:
    """Render selected figures and write the updated deck through the output policy."""

    result = update_figures_in_deck(presentation, figure_id=figure_id)
    write_patched_deck(
        presentation,
        result.deck,
        touched_slide_numbers=result.touched_slide_numbers,
        media_replacements=result.media_replacements,
        output=output,
    )
    return tuple(item.path for item in result.outputs)


def update_figures_in_deck(
    presentation: PresentationDefinition,
    *,
    figure_id: str | None = None,
    deck: PresentationObject | None = None,
) -> "FigureUpdateResult":
    """Render selected figures and replace anchors in an open deck."""

    active_deck = deck if deck is not None else Presentation(presentation.paths.deck_path)
    check_presentation(presentation, allow_shared_figure_relationships=True)
    ensure_generated_artifact_paths(presentation)
    anchors = discover_figure_anchors_in_deck(active_deck)
    selected_ids = _selected_figure_ids(presentation, figure_id, anchors)
    if not selected_ids:
        return FigureUpdateResult(active_deck, (), (), ())
    rendered = _run_selected_figures(presentation, selected_ids)
    bindings: list[tuple[str, pubify_data.BaseFigureResult, int, FigureAnchor]] = []
    for current_id in selected_ids:
        figure_result = rendered[current_id]
        for panel_index, anchor in _bind_anchors(current_id, figure_result, anchors):
            bindings.append((current_id, figure_result, panel_index, anchor))
    detach_keys = _shared_picture_detach_keys(
        all_anchors=anchors,
        selected_anchors=tuple(anchor for _current_id, _figure_result, _panel_index, anchor in bindings),
    )
    outputs: list[FigureOutput] = []
    touched_slide_numbers: list[int] = []
    media_replacements: list[MediaReplacement] = []
    _clear_selected_figure_pngs(presentation, selected_ids)

    for current_id, figure_result, panel_index, anchor in bindings:
        panel = figure_result.panels[panel_index - 1]
        render_options = _figure_render_options(figure_result, panel)
        output_path = _figure_output_path(presentation, current_id, len(figure_result.panels), panel_index)
        _render_panel_png(
            panel.payload,
            output_path,
            width_emu=anchor.shape.width,
            height_emu=anchor.shape.height,
            dpi=render_options.pop("dpi", presentation.config.defaults.dpi),
            preparation_options=render_options,
        )
        anchor_update = _replace_anchor_with_picture(
            anchor,
            output_path,
            force_distinct_media=_anchor_key(anchor) in detach_keys,
        )
        if anchor_update.touches_slide_xml:
            touched_slide_numbers.append(anchor.slide_number)
        if anchor_update.media_replacement is not None:
            media_replacements.append(anchor_update.media_replacement)
        outputs.append(FigureOutput(anchor.slide_number, anchor.shape_index, anchor.token, output_path))

    return FigureUpdateResult(
        active_deck,
        tuple(outputs),
        tuple(sorted(set(touched_slide_numbers))),
        tuple(media_replacements),
    )


def _selected_figure_ids(
    presentation: PresentationDefinition,
    figure_id: str | None,
    anchors: tuple[FigureAnchor, ...],
) -> tuple[str, ...]:
    available_ids = set(presentation.figures)
    if figure_id is None:
        return tuple(sorted({anchor.figure_id for anchor in anchors if anchor.figure_id in available_ids}))
    if figure_id not in available_ids:
        raise KeyError(f"Unknown figure '{figure_id}'")
    return (figure_id,)


def _run_selected_figures(
    presentation: PresentationDefinition,
    selected_ids: tuple[str, ...],
) -> dict[str, pubify_data.BaseFigureResult]:
    ctx = pubify_data.build_run_context(presentation.upstream)
    rendered: dict[str, pubify_data.BaseFigureResult] = {}
    for current_id in selected_ids:
        ((returned_id, result),) = pubify_data.run_figures(presentation.upstream, current_id, ctx=ctx)
        rendered[returned_id] = result
    return rendered


def _bind_anchors(
    figure_id: str,
    figure_result: pubify_data.BaseFigureResult,
    anchors: tuple[FigureAnchor, ...],
) -> tuple[tuple[int, FigureAnchor], ...]:
    figure_anchors = [anchor for anchor in anchors if anchor.figure_id == figure_id]
    panel_count = len(figure_result.panels)
    if panel_count == 1:
        scalar_anchors = [anchor for anchor in figure_anchors if anchor.panel_number is None]
        panel_anchors = [anchor for anchor in figure_anchors if anchor.panel_number == 1]
        extra_anchors = [anchor for anchor in figure_anchors if anchor.panel_number not in (None, 1)]
        if extra_anchors:
            tokens = ", ".join(anchor.token for anchor in extra_anchors)
            raise ValueError(f"Figure '{figure_id}' has extra panel anchors for a single-panel result: {tokens}")
        if len(scalar_anchors) + len(panel_anchors) != 1:
            raise ValueError(f"Figure '{figure_id}' requires exactly one anchor")
        return ((1, (scalar_anchors or panel_anchors)[0]),)

    bindings: list[tuple[int, FigureAnchor]] = []
    by_panel = {anchor.panel_number: anchor for anchor in figure_anchors if anchor.panel_number is not None}
    if any(anchor.panel_number is None for anchor in figure_anchors):
        raise ValueError(f"Figure '{figure_id}' returned multiple panels but has a scalar anchor")
    extra_panels = sorted(panel for panel in by_panel if panel is not None and panel > panel_count)
    if extra_panels:
        raise ValueError(f"Figure '{figure_id}' has extra panel anchors: {extra_panels}")
    for panel_index in sorted(panel for panel in by_panel if panel is not None):
        bindings.append((panel_index, by_panel[panel_index]))
    if not bindings:
        raise ValueError(f"Figure '{figure_id}' requires at least one panel anchor")
    return tuple(bindings)


def _shared_picture_detach_keys(
    *,
    all_anchors: tuple[FigureAnchor, ...],
    selected_anchors: tuple[FigureAnchor, ...],
) -> set[tuple[int, int]]:
    selected_keys = {_anchor_key(anchor) for anchor in selected_anchors}
    detach_keys: set[tuple[int, int]] = set()
    for group in _shared_picture_relationship_groups(all_anchors):
        group_keys = {_anchor_key(anchor) for anchor in group}
        group_selected_keys = group_keys & selected_keys
        if not group_selected_keys:
            continue
        if group_selected_keys == group_keys:
            keep = min(group, key=lambda anchor: anchor.shape_index)
            detach_keys.update(group_selected_keys - {_anchor_key(keep)})
        else:
            detach_keys.update(group_selected_keys)
    return detach_keys


def _shared_picture_relationship_groups(
    anchors: tuple[FigureAnchor, ...],
) -> tuple[tuple[FigureAnchor, ...], ...]:
    groups: dict[tuple[int, str], list[FigureAnchor]] = {}
    for anchor in anchors:
        r_ids = _embedded_image_r_ids(anchor.shape._element)
        if len(r_ids) == 1:
            groups.setdefault((anchor.slide_number, r_ids[0]), []).append(anchor)
    return tuple(
        tuple(group)
        for group in groups.values()
        if len(group) > 1 and len({anchor.token for anchor in group}) > 1
    )


def _anchor_key(anchor: FigureAnchor) -> tuple[int, int]:
    return (anchor.slide_number, anchor.shape_index)


def _clear_selected_figure_pngs(
    presentation: PresentationDefinition,
    selected_ids: tuple[str, ...],
) -> None:
    for current_id in selected_ids:
        for path in presentation.paths.figures_root.glob(f"{current_id}*.png"):
            if _is_figure_png_for_id(path, current_id):
                path.unlink()


def _is_figure_png_for_id(path: Path, figure_id: str) -> bool:
    stem = path.stem
    return stem == figure_id or re.fullmatch(rf"{re.escape(figure_id)}_[0-9]+", stem) is not None


def _figure_output_path(
    presentation: PresentationDefinition,
    figure_id: str,
    panel_count: int,
    panel_index: int,
) -> Path:
    if panel_count == 1:
        return presentation.paths.figures_root / f"{figure_id}.png"
    return presentation.paths.figures_root / f"{figure_id}_{panel_index}.png"


def _render_panel_png(
    payload: object,
    output_path: Path,
    *,
    width_emu: int,
    height_emu: int,
    dpi: int,
    preparation_options: dict[str, object] | None = None,
) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    render_options = dict(preparation_options or {})
    padding = _export_padding(render_options)
    if padding.has_side_specific:
        _save_fig_with_asymmetric_tight_padding(
            payload,
            output_path,
            width=width_emu / EMU_PER_INCH,
            height=height_emu / EMU_PER_INCH,
            dpi=dpi,
            padding=padding,
            render_options=render_options,
        )
        return

    pubify_mpl.save_fig(
        payload,
        output_path,
        width=width_emu / EMU_PER_INCH,
        height=height_emu / EMU_PER_INCH,
        **render_options,
        dpi=dpi,
        text_usetex=False,
        bbox_inches="tight",
        pad_inches=padding.symmetric,
    )


def _save_fig_with_asymmetric_tight_padding(
    payload: object,
    output_path: Path,
    *,
    width: float,
    height: float,
    dpi: int,
    padding: _ExportPadding,
    render_options: dict[str, object],
) -> None:
    with pubify_mpl.prepare_figure(
        payload,
        **render_options,
        dpi=dpi,
        text_usetex=False,
    ) as fig_export:
        fig_export.set_size_inches(width, height, forward=True)
        expanded_bbox = _asymmetric_tight_bbox(fig_export, padding)
        fig_export.savefig(
            output_path,
            dpi=dpi,
            bbox_inches=expanded_bbox,
            pad_inches=0.0,
        )


def _asymmetric_tight_bbox(figure: object, padding: _ExportPadding) -> Bbox:
    figure.canvas.draw()
    tight_bbox = figure.get_tightbbox(figure.canvas.get_renderer())
    return Bbox.from_extents(
        tight_bbox.x0 - padding.left,
        tight_bbox.y0 - padding.bottom,
        tight_bbox.x1 + padding.right,
        tight_bbox.y1 + padding.top,
    )


def _figure_render_options(
    figure_result: pubify_data.BaseFigureResult,
    panel: pubify_data.FigurePanel,
) -> dict[str, object]:
    options = dict(figure_result.metadata)
    options.update(panel.metadata)
    unknown = sorted(set(options) - FIGURE_PREPARATION_OPTIONS - {"dpi"})
    if unknown:
        joined = ", ".join(unknown)
        raise ValueError(f"Unsupported PowerPoint figure metadata option(s): {joined}")
    dpi = options.get("dpi")
    if dpi is not None and (isinstance(dpi, bool) or not isinstance(dpi, int) or dpi <= 0):
        raise ValueError("PowerPoint figure metadata option dpi must be a positive integer")
    _validate_export_padding_options(options)
    return options


def _validate_export_padding_options(options: Mapping[str, object]) -> None:
    for key in EXPORT_PADDING_OPTIONS:
        value = options.get(key)
        if value is None:
            continue
        if isinstance(value, bool) or not isinstance(value, Real) or not math.isfinite(float(value)) or value < 0:
            raise ValueError(f"PowerPoint figure metadata option {key} must be a non-negative number")


def _export_padding(options: dict[str, object]) -> _ExportPadding:
    symmetric = float(options.pop("export_pad_inches", 0.0))
    side_values: dict[str, float] = {}
    has_side_specific = False
    for side in ("left", "right", "top", "bottom"):
        key = f"export_pad_{side}_inches"
        if key in options:
            has_side_specific = True
            side_values[side] = float(options.pop(key))
        else:
            side_values[side] = symmetric
    return _ExportPadding(
        symmetric=symmetric,
        left=side_values["left"],
        right=side_values["right"],
        top=side_values["top"],
        bottom=side_values["bottom"],
        has_side_specific=has_side_specific,
    )


def _replace_anchor_with_picture(
    anchor: FigureAnchor,
    image_path: Path,
    *,
    force_distinct_media: bool = False,
) -> FigureAnchorUpdate:
    if force_distinct_media and _detach_picture_media(anchor, image_path):
        return FigureAnchorUpdate(media_replacement=None, touches_slide_xml=True)
    direct_update = _direct_picture_media_replacement(anchor, image_path)
    if direct_update is not None:
        return direct_update

    shape = anchor.shape
    slide = shape.part.slide
    left, top, width, height = shape.left, shape.top, shape.width, shape.height
    picture_left, picture_top, picture_width, picture_height = _contained_geometry(
        image_path,
        left=left,
        top=top,
        width=width,
        height=height,
    )
    parent = shape._element.getparent()
    for r_id in _embedded_image_r_ids(shape._element):
        shape.part.drop_rel(r_id)
    parent.remove(shape._element)
    picture = slide.shapes.add_picture(
        str(image_path),
        picture_left,
        picture_top,
        width=picture_width,
        height=picture_height,
    )
    set_shape_alt_text(picture, anchor.token)
    return FigureAnchorUpdate(media_replacement=None, touches_slide_xml=True)


def _detach_picture_media(anchor: FigureAnchor, image_path: Path) -> bool:
    r_ids = _embedded_image_r_ids(anchor.shape._element)
    if len(r_ids) != 1:
        return False
    image_part, _reused_r_id = anchor.shape.part.get_or_add_image_part(str(image_path))
    r_id = anchor.shape.part.rels._add_relationship(RT.IMAGE, image_part, is_external=False)
    _set_embedded_image_r_id(anchor.shape._element, r_ids[0], r_id)
    return True


def _direct_picture_media_replacement(anchor: FigureAnchor, image_path: Path) -> FigureAnchorUpdate | None:
    r_ids = _embedded_image_r_ids(anchor.shape._element)
    if len(r_ids) != 1:
        return None
    relationship = anchor.shape.part.rels[r_ids[0]]
    target_part = relationship.target_part
    part_name = str(target_part.partname)
    if target_part.content_type != "image/png" or not part_name.lower().endswith(".png"):
        return None
    touches_slide_xml = False
    if not _aspect_ratio_matches_picture_geometry(anchor.shape, image_path):
        _set_picture_geometry_to_contained_image(anchor.shape, image_path)
        touches_slide_xml = True
    return FigureAnchorUpdate(
        media_replacement=MediaReplacement(part_name=part_name, blob=image_path.read_bytes()),
        touches_slide_xml=touches_slide_xml,
    )


def _aspect_ratio_matches_picture_geometry(shape: object, image_path: Path) -> bool:
    with Image.open(image_path) as image:
        image_aspect = image.width / image.height
    geometry_aspect = shape.width / shape.height
    return abs(image_aspect - geometry_aspect) / geometry_aspect <= ASPECT_RATIO_TOLERANCE


def _set_picture_geometry_to_contained_image(shape: object, image_path: Path) -> None:
    left, top, width, height = _contained_geometry(
        image_path,
        left=shape.left,
        top=shape.top,
        width=shape.width,
        height=shape.height,
    )
    xfrm_matches = shape._element.xpath(".//a:xfrm")
    if not xfrm_matches:
        raise ValueError("Picture shape does not expose a transform node")
    xfrm = xfrm_matches[0]
    off_matches = xfrm.xpath("./a:off")
    ext_matches = xfrm.xpath("./a:ext")
    if not off_matches or not ext_matches:
        raise ValueError("Picture shape does not expose transform geometry")
    off_matches[0].set("x", str(int(left)))
    off_matches[0].set("y", str(int(top)))
    ext_matches[0].set("cx", str(int(width)))
    ext_matches[0].set("cy", str(int(height)))


def _embedded_image_r_ids(element: object) -> tuple[str, ...]:
    return tuple(
        r_id
        for blip in element.xpath(".//a:blip")
        if (r_id := blip.get(qn("r:embed"))) is not None
    )


def _set_embedded_image_r_id(element: object, old_r_id: str, new_r_id: str) -> None:
    for blip in element.xpath(".//a:blip"):
        if blip.get(qn("r:embed")) == old_r_id:
            blip.set(qn("r:embed"), new_r_id)
            return


def _contained_geometry(
    image_path: Path,
    *,
    left: int,
    top: int,
    width: int,
    height: int,
) -> tuple[Emu, Emu, Emu, Emu]:
    with Image.open(image_path) as image:
        image_width, image_height = image.size
    image_aspect = image_width / image_height
    box_aspect = width / height
    if image_aspect >= box_aspect:
        picture_width = width
        picture_height = round(width / image_aspect)
    else:
        picture_height = height
        picture_width = round(height * image_aspect)
    picture_left = left + round((width - picture_width) / 2)
    picture_top = top + round((height - picture_height) / 2)
    return Emu(picture_left), Emu(picture_top), Emu(picture_width), Emu(picture_height)
