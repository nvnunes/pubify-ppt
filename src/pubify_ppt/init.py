from __future__ import annotations

from collections.abc import Callable
from importlib import resources
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches

from pubify_ppt.config import (
    PRESENTATION_CONFIG_FILENAME,
    WORKSPACE_CONFIG_FILENAME,
    load_workspace_config,
    render_default_workspace_config,
    workspace_config_has_section,
    write_default_presentation_config,
    write_default_workspace_config,
)


PPT_ARTIFACTS_NAMESPACE = "ppt-artifacts"
FIGURES_ARTIFACT_DIR = "figures"
BACKUPS_ARTIFACT_DIR = "backups"
PRESENTATION_ENTRYPOINT = "figures.py"
PRESENTATIONS_AGENTS_FILENAME = "AGENTS.md"
STARTER_DATA_FILENAME = "example.csv"
STARTER_FIGURE_TOKEN = "{{fig:example}}"
STARTER_STAT_TOKEN = "{{stat:example.count}}"
STARTER_TABLE_TOKEN = "{{table:example}}"


def init_workspace(workspace_root: Path) -> Path:
    """Create or update the workspace-level ``pubify-ppt`` scaffold."""

    root = workspace_root.resolve()
    config_path = root / WORKSPACE_CONFIG_FILENAME
    if not config_path.exists():
        write_default_workspace_config(config_path)
    elif not workspace_config_has_section(root):
        _append_workspace_section(config_path)

    workspace = load_workspace_config(root)
    workspace.presentations_root.mkdir(parents=True, exist_ok=True)
    _write_if_missing(
        workspace.presentations_root / PRESENTATIONS_AGENTS_FILENAME,
        write_presentations_agents_file,
    )
    return root


def init_presentation_by_id(workspace_root: Path, presentation_id: str) -> Path:
    """Create a presentation scaffold without overwriting user-owned files."""

    if presentation_id in {"", ".", ".."} or "/" in presentation_id or "\\" in presentation_id:
        raise ValueError("presentation id must be a single non-empty path segment")

    workspace = load_workspace_config(workspace_root)
    _write_if_missing(
        workspace.presentations_root / PRESENTATIONS_AGENTS_FILENAME,
        write_presentations_agents_file,
    )
    presentation_root = workspace.presentations_root / presentation_id
    data_root = presentation_root / "data"
    artifacts_root = data_root / PPT_ARTIFACTS_NAMESPACE

    presentation_root.mkdir(parents=True, exist_ok=True)
    _ensure_data_root(data_root)
    (artifacts_root / FIGURES_ARTIFACT_DIR).mkdir(parents=True, exist_ok=True)
    (artifacts_root / BACKUPS_ARTIFACT_DIR).mkdir(parents=True, exist_ok=True)

    _write_if_missing(presentation_root / PRESENTATION_CONFIG_FILENAME, write_default_presentation_config)
    _write_if_missing(presentation_root / PRESENTATION_ENTRYPOINT, write_starter_figures_module)
    _write_text_if_missing(data_root / STARTER_DATA_FILENAME, "x,y\n1,1\n2,4\n3,9\n")
    _write_if_missing(presentation_root / "deck.pptx", write_starter_deck)
    return presentation_root


def write_starter_figures_module(path: Path) -> None:
    """Write a runnable starter ``figures.py`` module."""

    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(_render_starter_figures_module(), encoding="utf-8")


def write_presentations_agents_file(path: Path) -> None:
    """Write the shared presentations-root ``AGENTS.md`` scaffold for ``ppt init``."""

    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(_load_init_asset_text("AGENTS.example.md"), encoding="utf-8")


def write_starter_deck(path: Path) -> None:
    """Create the starter editable source deck with figure, stat, and table anchors."""

    path.parent.mkdir(parents=True, exist_ok=True)
    deck = Presentation()
    blank_layout = deck.slide_layouts[6]
    slide = deck.slides.add_slide(blank_layout)

    title_box = slide.shapes.add_textbox(Inches(0.6), Inches(0.25), Inches(8.0), Inches(0.5))
    title_box.text = "pubify-ppt starter"

    figure_shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0.8),
        Inches(1.0),
        Inches(5.6),
        Inches(3.5),
    )
    figure_shape.text = "Figure: {{fig:example}}"
    _set_shape_alt_text(figure_shape, STARTER_FIGURE_TOKEN)

    stat_box = slide.shapes.add_textbox(Inches(0.8), Inches(4.85), Inches(5.6), Inches(0.6))
    stat_box.text = f"Example count: {STARTER_STAT_TOKEN}"

    table_label = slide.shapes.add_textbox(Inches(6.75), Inches(1.0), Inches(2.4), Inches(0.35))
    table_label.text = "Example table"

    table_shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(6.75),
        Inches(1.45),
        Inches(2.25),
        Inches(1.35),
    )
    table_shape.text = STARTER_TABLE_TOKEN

    deck.save(path)


def _append_workspace_section(config_path: Path) -> None:
    current = config_path.read_text(encoding="utf-8")
    separator = "" if not current or current.endswith("\n") else "\n"
    spacer = "" if not current.strip() else "\n"
    config_path.write_text(
        f"{current}{separator}{spacer}{render_default_workspace_config()}",
        encoding="utf-8",
    )


def _ensure_data_root(data_root: Path) -> None:
    if data_root.exists() or data_root.is_symlink():
        if not data_root.is_dir():
            raise ValueError(
                "Presentation data path must be a directory or symlink "
                f"to a directory: {data_root}"
            )
        return
    data_root.mkdir(parents=True, exist_ok=True)


def _write_if_missing(path: Path, writer: Callable[[Path], None]) -> None:
    if path.exists():
        return
    writer(path)


def _write_text_if_missing(path: Path, text: str) -> None:
    if path.exists():
        return
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(text, encoding="utf-8")


def _set_shape_alt_text(shape: object, value: str) -> None:
    c_nv_pr = shape._element.xpath(".//p:cNvPr")[0]
    c_nv_pr.set("descr", value)


def _load_init_asset_text(filename: str) -> str:
    return resources.files("pubify_ppt.assets.init").joinpath(filename).read_text(encoding="utf-8")


def _render_starter_figures_module() -> str:
    return '''"""Figures entrypoint for a pubify-ppt presentation."""

import csv

import matplotlib.pyplot as plt
import numpy as np

from pubify_data import data, figure, stat, table
from pubify_ppt import FigureResult, StatResult, TableResult


@data("example.csv")
def load_example(ctx, path):
    with path.open(newline="", encoding="utf-8") as handle:
        rows = list(csv.DictReader(handle))
    return {
        "x": np.array([float(row["x"]) for row in rows]),
        "y": np.array([float(row["y"]) for row in rows]),
    }


@figure
def plot_example(ctx, example):
    fig, ax = plt.subplots()
    ax.plot(example["x"], example["y"], marker="o")
    ax.set_xlabel("x")
    ax.set_ylabel("y")
    ax.set_title("Starter figure")
    return FigureResult(fig)


@stat
def compute_example(ctx, example):
    return StatResult({"count": example["x"].size})


@table
def tabulate_example(ctx, example):
    rows = list(zip(example["x"].astype(int), example["y"].astype(int)))
    return TableResult(rows, metadata={"columns": ("x", "y")})
'''
