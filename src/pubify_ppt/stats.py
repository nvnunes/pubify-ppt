from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
import re

from pptx import Presentation
from pptx.presentation import Presentation as PresentationObject
import pubify_data

from pubify_ppt.anchors import STAT_RENDERED_TOKEN_RE, STAT_TOKEN_RE, set_shape_alt_text, shape_alt_text, split_stat_reference
from pubify_ppt.backups import write_deck
from pubify_ppt.discovery import PresentationDefinition
from pubify_ppt.runtime import check_presentation


@dataclass(frozen=True)
class StatReplacement:
    """One stat token replacement applied to a deck."""

    slide_number: int
    shape_index: int
    stat_id: str
    key: str | None
    token: str
    value: str


@dataclass(frozen=True)
class StatUpdateResult:
    """Stat replacements applied to an open deck."""

    deck: PresentationObject
    replacements: tuple[StatReplacement, ...]


@dataclass(frozen=True)
class RenderedStatAnchor:
    """One text shape carrying a rendered stat marker in Alt Text."""

    stat_id: str
    key: str | None
    token: str
    old_value: str


def update_stats(
    presentation: PresentationDefinition,
    *,
    stat_id: str | None = None,
) -> tuple[StatReplacement, ...]:
    """Compute selected stats and replace matching inline PowerPoint tokens."""

    result = update_stats_in_deck(presentation, stat_id=stat_id)
    write_deck(presentation, result.deck)
    return result.replacements


def update_stats_to_output(
    presentation: PresentationDefinition,
    *,
    stat_id: str | None = None,
    output: Path | None = None,
) -> tuple[StatReplacement, ...]:
    """Compute selected stats and write the updated deck through the output policy."""

    result = update_stats_in_deck(presentation, stat_id=stat_id)
    write_deck(presentation, result.deck, output=output)
    return result.replacements


def update_stats_in_deck(
    presentation: PresentationDefinition,
    *,
    stat_id: str | None = None,
    deck: PresentationObject | None = None,
) -> StatUpdateResult:
    """Compute selected stats and replace matching tokens in an open deck."""

    check_presentation(presentation)
    selected_ids = _selected_stat_ids(presentation, stat_id)
    active_deck = deck if deck is not None else Presentation(presentation.paths.deck_path)
    if not selected_ids:
        return StatUpdateResult(active_deck, ())

    values = _run_selected_stats(presentation, selected_ids)
    replacements: list[StatReplacement] = []

    for slide_number, slide in enumerate(active_deck.slides, start=1):
        for shape_index, shape in enumerate(slide.shapes, start=1):
            if getattr(shape, "text_frame", None) is None:
                continue
            anchor = _parse_rendered_stat_anchor(shape_alt_text(shape), stat_ids=selected_ids)
            if anchor is not None:
                if anchor.stat_id in selected_ids:
                    replacements.append(
                        _refresh_anchored_stat(
                            shape,
                            slide_number=slide_number,
                            shape_index=shape_index,
                            anchor=anchor,
                            values=values,
                        )
                    )
                continue
            replacements.extend(
                _replace_new_stat_token(
                    shape,
                    slide_number=slide_number,
                    shape_index=shape_index,
                    selected_ids=selected_ids,
                    values=values,
                )
            )

    return StatUpdateResult(active_deck, tuple(replacements))


def _selected_stat_ids(presentation: PresentationDefinition, stat_id: str | None) -> tuple[str, ...]:
    available_ids = set(presentation.stats)
    if stat_id is None:
        return tuple(sorted(available_ids))
    if stat_id not in available_ids:
        raise KeyError(f"Unknown stat '{stat_id}'")
    return (stat_id,)


def _run_selected_stats(
    presentation: PresentationDefinition,
    selected_ids: tuple[str, ...],
) -> dict[tuple[str, str | None], str]:
    ctx = pubify_data.build_run_context(presentation.upstream)
    values: dict[tuple[str, str | None], str] = {}
    for current_id in selected_ids:
        (computed,) = pubify_data.run_stats(presentation.upstream, current_id, ctx=ctx)
        for value in computed.values:
            values[(computed.stat_id, value.key)] = value.value
    return values


def _replace_new_stat_token(
    shape: object,
    *,
    slide_number: int,
    shape_index: int,
    selected_ids: tuple[str, ...],
    values: dict[tuple[str, str | None], str],
) -> tuple[StatReplacement, ...]:
    matches = _shape_stat_matches(shape)
    parsed = [(match, *split_stat_reference(match.group(1), stat_ids=selected_ids)) for match in matches]
    selected = [item for item in parsed if item[1] in selected_ids]
    if not selected:
        return ()
    if len(matches) != 1:
        raise ValueError(
            f"Slide {slide_number}: stat-managed text boxes must contain exactly one stat token; "
            "split multiple stats into separate text boxes"
        )
    match, stat_id, key = selected[0]
    token = match.group(0)
    value = _stat_value(stat_id, key, values)
    _replace_shape_text_once(
        shape,
        token,
        value,
        slide_number=slide_number,
        missing_message=(
            f"Slide {slide_number}: stat token {token} was discovered but could not be replaced; "
            "retype the token in one text box"
        ),
    )
    set_shape_alt_text(shape, _rendered_stat_anchor(token, value))
    return (StatReplacement(slide_number, shape_index, stat_id, key, token, value),)


def _refresh_anchored_stat(
    shape: object,
    *,
    slide_number: int,
    shape_index: int,
    anchor: RenderedStatAnchor,
    values: dict[tuple[str, str | None], str],
) -> StatReplacement:
    value = _stat_value(anchor.stat_id, anchor.key, values)
    if anchor.old_value and _shape_text_count(shape, anchor.old_value) == 1:
        _replace_shape_text_once(
            shape,
            anchor.old_value,
            value,
            slide_number=slide_number,
            missing_message=_repair_message(slide_number, anchor),
        )
    elif _shape_text_count(shape, anchor.token) == 1:
        _replace_shape_text_once(
            shape,
            anchor.token,
            value,
            slide_number=slide_number,
            missing_message=_repair_message(slide_number, anchor),
        )
    else:
        raise ValueError(_repair_message(slide_number, anchor))
    set_shape_alt_text(shape, _rendered_stat_anchor(anchor.token, value))
    return StatReplacement(slide_number, shape_index, anchor.stat_id, anchor.key, anchor.token, value)


def _parse_rendered_stat_anchor(alt_text: str | None, *, stat_ids: tuple[str, ...]) -> RenderedStatAnchor | None:
    if not alt_text:
        return None
    match = STAT_RENDERED_TOKEN_RE.fullmatch(alt_text.strip())
    if match is None:
        return None
    stat_id, key = split_stat_reference(match.group(1), stat_ids=stat_ids)
    token = _stat_token(stat_id, key)
    return RenderedStatAnchor(stat_id=stat_id, key=key, token=token, old_value=match.group(2))


def _stat_token(stat_id: str, key: str | None) -> str:
    return f"{{{{stat:{stat_id}{'.' + key if key is not None else ''}}}}}"


def _rendered_stat_anchor(token: str, value: str) -> str:
    inner = token.removeprefix("{{stat:").removesuffix("}}")
    return f"{{{{stat:{inner}={value}}}}}"


def _stat_value(stat_id: str, key: str | None, values: dict[tuple[str, str | None], str]) -> str:
    value_key = (stat_id, key)
    if value_key not in values:
        if key is None:
            raise ValueError(f"Missing scalar value for stat '{stat_id}'")
        raise ValueError(f"Missing key '{key}' for stat '{stat_id}'")
    return values[value_key]


def _shape_stat_matches(shape: object) -> list[re.Match[str]]:
    matches: list[re.Match[str]] = []
    text_frame = getattr(shape, "text_frame", None)
    if text_frame is None:
        return matches
    for paragraph in text_frame.paragraphs:
        paragraph_text = "".join(run.text for run in paragraph.runs)
        matches.extend(STAT_TOKEN_RE.finditer(paragraph_text))
    return matches


def _shape_text_count(shape: object, text: str) -> int:
    if text == "":
        return 0
    text_frame = getattr(shape, "text_frame", None)
    if text_frame is None:
        return 0
    return sum("".join(run.text for run in paragraph.runs).count(text) for paragraph in text_frame.paragraphs)


def _replace_shape_text_once(
    shape: object,
    old: str,
    new: str,
    *,
    slide_number: int,
    missing_message: str,
) -> None:
    if old == "":
        raise ValueError(missing_message)
    text_frame = getattr(shape, "text_frame", None)
    if text_frame is None:
        raise ValueError(missing_message)
    matches: list[object] = []
    for paragraph in text_frame.paragraphs:
        paragraph_text = "".join(run.text for run in paragraph.runs)
        count = paragraph_text.count(old)
        if count:
            matches.extend([paragraph] * count)
    if len(matches) != 1:
        if len(matches) > 1:
            raise ValueError(
                f"Slide {slide_number}: replacement text {old!r} appears more than once in a stat-managed text box; "
                "restore the stat token so pubify-ppt can update it safely"
            )
        raise ValueError(missing_message)
    paragraph = matches[0]
    _replace_paragraph_text_once(paragraph, old, new)


def _replace_paragraph_text_once(paragraph: object, old: str, new: str) -> None:
    runs = list(paragraph.runs)
    run_texts = [run.text for run in runs]
    paragraph_text = "".join(run_texts)
    start = paragraph_text.index(old)
    end = start + len(old)
    rendered = _render_replaced_runs(run_texts, start=start, end=end, replacement=new)
    for index, run in enumerate(runs):
        run.text = rendered[index]


def _render_replaced_runs(run_texts: list[str], *, start: int, end: int, replacement: str) -> list[str]:
    rendered = [""] * len(run_texts)
    position = 0
    inserted = False
    for index, text in enumerate(run_texts):
        run_start = position
        run_end = position + len(text)
        if run_start < start:
            rendered[index] += text[: min(start, run_end) - run_start]
        if not inserted and run_start <= start < run_end:
            rendered[index] += replacement
            inserted = True
        if run_end > end:
            rendered[index] += text[max(end, run_start) - run_start :]
        position = run_end
    if not inserted and rendered:
        rendered[-1] += replacement
    return rendered


def _repair_message(slide_number: int, anchor: RenderedStatAnchor) -> str:
    return (
        f"Slide {slide_number}: can't update {anchor.token}; previous value {anchor.old_value!r} "
        f"was not found exactly once. Restore it or reinsert {anchor.token}."
    )
