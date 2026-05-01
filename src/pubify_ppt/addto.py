from __future__ import annotations

from dataclasses import dataclass

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.presentation import Presentation as PresentationObject
import pubify_data

from pubify_ppt.anchors import (
    FIGURE_TOKEN_RE,
    STAT_TOKEN_RE,
    TABLE_TOKEN_RE,
    set_shape_alt_text,
    split_stat_reference,
)
from pubify_ppt.backups import write_patched_deck
from pubify_ppt.discovery import PresentationDefinition
from pubify_ppt.figures import FigureOutput, update_figures_in_deck
from pubify_ppt.stats import StatReplacement, update_stats_in_deck
from pubify_ppt.tables import TableReplacement, update_tables_in_deck


FIGURE_BOX_WIDTH_FRACTION = 0.60
FIGURE_BOX_HEIGHT_FRACTION = 0.45
STAT_BOX_WIDTH_FRACTION = 0.45
STAT_BOX_HEIGHT_FRACTION = 0.08
TABLE_BOX_WIDTH_FRACTION = 0.45
TABLE_BOX_HEIGHT_FRACTION = 0.25


@dataclass(frozen=True)
class AddedAnchor:
    """One managed PowerPoint anchor inserted into a deck slide."""

    slide_number: int
    token: str


def add_figure_anchor(
    presentation: PresentationDefinition,
    *,
    figure_ref: str,
    slide_number: int,
) -> tuple[FigureOutput, ...]:
    """Add a centered figure anchor and render the selected figure into the deck."""

    deck = Presentation(presentation.paths.deck_path)
    _slide_by_number(deck, slide_number)
    token, figure_id, rendered = _rendered_figure_token(presentation, figure_ref)
    _add_figure_anchor_token_to_deck(deck=deck, token=token, slide_number=slide_number)
    result = update_figures_in_deck(
        presentation,
        figure_id=figure_id,
        deck=deck,
        _rendered_figures=rendered,
    )
    write_patched_deck(
        presentation,
        result.deck,
        touched_slide_numbers=result.touched_slide_numbers,
        media_replacements=result.media_replacements,
    )
    return result.outputs


def add_stat_anchor(
    presentation: PresentationDefinition,
    *,
    stat_ref: str,
    slide_number: int,
) -> tuple[StatReplacement, ...]:
    """Add a centered stat token and render the selected stat into the deck."""

    deck = Presentation(presentation.paths.deck_path)
    _slide_by_number(deck, slide_number)
    _token, stat_id = _stat_token(presentation, stat_ref)
    add_stat_anchor_to_deck(presentation, deck=deck, stat_ref=stat_ref, slide_number=slide_number)
    result = update_stats_in_deck(presentation, stat_id=stat_id, deck=deck)
    write_patched_deck(presentation, result.deck, touched_slide_numbers=_replacement_slide_numbers(result.replacements))
    return result.replacements


def add_table_anchor(
    presentation: PresentationDefinition,
    *,
    table_id: str,
    slide_number: int,
) -> tuple[TableReplacement, ...]:
    """Add a centered table anchor and render the selected table into the deck."""

    deck = Presentation(presentation.paths.deck_path)
    _slide_by_number(deck, slide_number)
    add_table_anchor_to_deck(presentation, deck=deck, table_id=table_id, slide_number=slide_number)
    result = update_tables_in_deck(presentation, table_id=table_id, deck=deck)
    write_patched_deck(presentation, result.deck, touched_slide_numbers=_replacement_slide_numbers(result.replacements))
    return result.replacements


def add_figure_anchor_to_deck(
    presentation: PresentationDefinition,
    *,
    deck: PresentationObject,
    figure_ref: str,
    slide_number: int,
) -> AddedAnchor:
    """Add a figure anchor to an open deck without writing it."""

    token, _figure_id = _figure_token(presentation, figure_ref)
    return _add_figure_anchor_token_to_deck(deck=deck, token=token, slide_number=slide_number)


def _add_figure_anchor_token_to_deck(
    *,
    deck: PresentationObject,
    token: str,
    slide_number: int,
) -> AddedAnchor:
    slide = _slide_by_number(deck, slide_number)
    left, top, width, height = _centered_geometry(
        deck,
        width_fraction=FIGURE_BOX_WIDTH_FRACTION,
        height_fraction=FIGURE_BOX_HEIGHT_FRACTION,
    )
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.text = token
    set_shape_alt_text(shape, token)
    return AddedAnchor(slide_number=slide_number, token=token)


def add_stat_anchor_to_deck(
    presentation: PresentationDefinition,
    *,
    deck: PresentationObject,
    stat_ref: str,
    slide_number: int,
) -> AddedAnchor:
    """Add a stat token to an open deck without writing it."""

    token, _stat_id = _stat_token(presentation, stat_ref)
    slide = _slide_by_number(deck, slide_number)
    left, top, width, height = _centered_geometry(
        deck,
        width_fraction=STAT_BOX_WIDTH_FRACTION,
        height_fraction=STAT_BOX_HEIGHT_FRACTION,
    )
    shape = slide.shapes.add_textbox(left, top, width, height)
    shape.text = token
    return AddedAnchor(slide_number=slide_number, token=token)


def add_table_anchor_to_deck(
    presentation: PresentationDefinition,
    *,
    deck: PresentationObject,
    table_id: str,
    slide_number: int,
) -> AddedAnchor:
    """Add a table anchor to an open deck without writing it."""

    token, _table_id = _table_token(presentation, table_id)
    slide = _slide_by_number(deck, slide_number)
    left, top, width, height = _centered_geometry(
        deck,
        width_fraction=TABLE_BOX_WIDTH_FRACTION,
        height_fraction=TABLE_BOX_HEIGHT_FRACTION,
    )
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    set_shape_alt_text(shape, token)
    return AddedAnchor(slide_number=slide_number, token=token)


def _figure_token(presentation: PresentationDefinition, figure_ref: str) -> tuple[str, str]:
    token = f"{{{{fig:{figure_ref}}}}}"
    match = FIGURE_TOKEN_RE.fullmatch(token)
    if match is None:
        raise ValueError(f"Invalid figure reference '{figure_ref}'")
    figure_id = match.group(1)
    if figure_id not in presentation.figures:
        raise KeyError(f"Unknown figure '{figure_id}'")
    return match.group(0), figure_id


def _rendered_figure_token(
    presentation: PresentationDefinition,
    figure_ref: str,
) -> tuple[str, str, dict[str, pubify_data.BaseFigureResult]]:
    token, figure_id = _figure_token(presentation, figure_ref)
    rendered_figures = _run_one_figure(presentation, figure_id)
    rendered = rendered_figures[figure_id]
    match = FIGURE_TOKEN_RE.fullmatch(token)
    if match is None:
        raise ValueError(f"Invalid figure reference '{figure_ref}'")
    panel_number = int(match.group(2)) if match.group(2) is not None else None
    panel_count = len(rendered.panels)
    if panel_number is None and panel_count > 1:
        return f"{{{{fig:{figure_id}:1}}}}", figure_id, rendered_figures
    if panel_number is not None and panel_number > panel_count:
        raise ValueError(f"Figure '{figure_id}' has {panel_count} panel(s), cannot add panel {panel_number}")
    return token, figure_id, rendered_figures


def _run_one_figure(
    presentation: PresentationDefinition,
    figure_id: str,
) -> dict[str, pubify_data.BaseFigureResult]:
    ctx = pubify_data.build_run_context(presentation.upstream)
    ((returned_id, result),) = pubify_data.run_figures(presentation.upstream, figure_id, ctx=ctx)
    return {returned_id: result}


def _stat_token(presentation: PresentationDefinition, stat_ref: str) -> tuple[str, str]:
    token = f"{{{{stat:{stat_ref}}}}}"
    match = STAT_TOKEN_RE.fullmatch(token)
    if match is None:
        raise ValueError(f"Invalid stat reference '{stat_ref}'")
    stat_id, _key = split_stat_reference(match.group(1), stat_ids=set(presentation.stats))
    if stat_id not in presentation.stats:
        raise KeyError(f"Unknown stat '{stat_id}'")
    return match.group(0), stat_id


def _table_token(presentation: PresentationDefinition, table_id: str) -> tuple[str, str]:
    token = f"{{{{table:{table_id}}}}}"
    match = TABLE_TOKEN_RE.fullmatch(token)
    if match is None:
        raise ValueError(f"Invalid table id '{table_id}'")
    if table_id not in presentation.tables:
        raise KeyError(f"Unknown table '{table_id}'")
    return match.group(0), table_id


def _slide_by_number(deck: PresentationObject, slide_number: int) -> object:
    if slide_number < 1 or slide_number > len(deck.slides):
        raise ValueError(f"Slide number must be between 1 and {len(deck.slides)}: {slide_number}")
    return deck.slides[slide_number - 1]


def _centered_geometry(
    deck: PresentationObject,
    *,
    width_fraction: float,
    height_fraction: float,
) -> tuple[int, int, int, int]:
    width = round(int(deck.slide_width) * width_fraction)
    height = round(int(deck.slide_height) * height_fraction)
    left = round((int(deck.slide_width) - width) / 2)
    top = round((int(deck.slide_height) - height) / 2)
    return left, top, width, height


def _replacement_slide_numbers(
    replacements: tuple[StatReplacement, ...] | tuple[TableReplacement, ...],
) -> tuple[int, ...]:
    return tuple(sorted({replacement.slide_number for replacement in replacements}))
