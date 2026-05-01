from __future__ import annotations

from collections.abc import Sequence
from dataclasses import dataclass
from pathlib import Path

from pptx import Presentation
from pptx.presentation import Presentation as PresentationObject
import pubify_data

from pubify_ppt.anchors import TableAnchor, discover_table_anchors_in_deck, set_shape_alt_text
from pubify_ppt.backups import write_patched_deck
from pubify_ppt.discovery import PresentationDefinition
from pubify_ppt.runtime import check_presentation


TABLE_COLUMNS_METADATA_KEY = "columns"


@dataclass(frozen=True)
class RenderedTable:
    """Text-only table payload ready to write into PowerPoint."""

    table_id: str
    width: int
    metadata: dict[str, object]
    body: tuple[tuple[str, ...], ...]

    @property
    def row_count(self) -> int:
        return 1 + len(self.body)

    @property
    def column_count(self) -> int:
        return self.width


@dataclass(frozen=True)
class TableReplacement:
    """One table replacement applied to a deck."""

    slide_number: int
    shape_index: int
    table_id: str
    token: str
    rows: int
    columns: int

    @property
    def summary(self) -> str:
        return f"{self.rows} rows x {self.columns} columns"


@dataclass(frozen=True)
class TableUpdateResult:
    """Table replacements applied to an open deck."""

    deck: PresentationObject
    replacements: tuple[TableReplacement, ...]


def update_tables(
    presentation: PresentationDefinition,
    *,
    table_id: str | None = None,
) -> tuple[TableReplacement, ...]:
    """Compute selected tables and update matching PowerPoint table anchors."""

    result = update_tables_in_deck(presentation, table_id=table_id)
    write_patched_deck(presentation, result.deck, touched_slide_numbers=_touched_slides(result.replacements))
    return result.replacements


def update_tables_to_output(
    presentation: PresentationDefinition,
    *,
    table_id: str | None = None,
    output: Path | None = None,
) -> tuple[TableReplacement, ...]:
    """Compute selected tables and write the updated deck through the output policy."""

    result = update_tables_in_deck(presentation, table_id=table_id)
    write_patched_deck(
        presentation,
        result.deck,
        touched_slide_numbers=_touched_slides(result.replacements),
        output=output,
    )
    return result.replacements


def update_tables_in_deck(
    presentation: PresentationDefinition,
    *,
    table_id: str | None = None,
    deck: PresentationObject | None = None,
) -> TableUpdateResult:
    """Compute selected tables and update native PowerPoint tables in an open deck."""

    check_presentation(presentation, allow_shared_figure_relationships=True)
    active_deck = deck if deck is not None else Presentation(presentation.paths.deck_path)
    anchors = discover_table_anchors_in_deck(active_deck)
    selected_ids = _selected_table_ids(presentation, table_id, anchors)
    if not selected_ids:
        return TableUpdateResult(active_deck, ())

    rendered = _run_selected_tables(presentation, selected_ids)
    replacements: list[TableReplacement] = []
    anchors_by_id: dict[str, list[TableAnchor]] = {}
    for anchor in anchors:
        anchors_by_id.setdefault(anchor.table_id, []).append(anchor)

    for current_id in selected_ids:
        current_anchors = anchors_by_id.get(current_id, [])
        if not current_anchors:
            raise ValueError(f"Table '{current_id}' requires at least one anchor")
        table = rendered[current_id]
        for anchor in current_anchors:
            _apply_table(anchor, table)
            replacements.append(
                TableReplacement(
                    slide_number=anchor.slide_number,
                    shape_index=anchor.shape_index,
                    table_id=current_id,
                    token=anchor.token,
                    rows=table.row_count,
                    columns=table.column_count,
                )
            )

    return TableUpdateResult(active_deck, tuple(replacements))


def _selected_table_ids(
    presentation: PresentationDefinition,
    table_id: str | None,
    anchors: tuple[TableAnchor, ...],
) -> tuple[str, ...]:
    available_ids = set(presentation.tables)
    if table_id is None:
        return tuple(sorted({anchor.table_id for anchor in anchors if anchor.table_id in available_ids}))
    if table_id not in available_ids:
        raise KeyError(f"Unknown table '{table_id}'")
    return (table_id,)


def _touched_slides(replacements: tuple[TableReplacement, ...]) -> tuple[int, ...]:
    return tuple(sorted({replacement.slide_number for replacement in replacements}))


def _run_selected_tables(
    presentation: PresentationDefinition,
    selected_ids: tuple[str, ...],
) -> dict[str, RenderedTable]:
    ctx = pubify_data.build_run_context(presentation.upstream)
    rendered: dict[str, RenderedTable] = {}
    for current_id in selected_ids:
        (computed,) = pubify_data.run_tables(presentation.upstream, current_id, ctx=ctx)
        rendered[computed.table_id] = _render_table(computed)
    return rendered


def _render_table(computed: pubify_data.ComputedTable) -> RenderedTable:
    if len(computed.bodies) != 1:
        raise ValueError(
            f"Table '{computed.table_id}' returned {len(computed.bodies)} bodies; "
            "PowerPoint table anchors support one body in v1"
        )
    body = tuple(tuple(_cell_text(cell) for cell in row) for row in computed.bodies[0])
    return RenderedTable(table_id=computed.table_id, width=computed.width, metadata=computed.metadata, body=body)


def _default_columns(table: RenderedTable) -> tuple[str, ...]:
    value = table.metadata.get(TABLE_COLUMNS_METADATA_KEY)
    if value is None:
        return tuple("" for _ in range(table.width))
    if not _is_sequence(value):
        raise ValueError(f"Table '{table.table_id}' metadata['columns'] must be a sequence of strings")
    columns = tuple(value)
    if not all(isinstance(column, str) for column in columns):
        raise ValueError(f"Table '{table.table_id}' metadata['columns'] must be a sequence of strings")
    if len(columns) != table.width:
        raise ValueError(
            f"Table '{table.table_id}' metadata['columns'] has {len(columns)} columns but data has {table.width}"
        )
    return columns


def _is_sequence(value: object) -> bool:
    return isinstance(value, Sequence) and not isinstance(value, (str, bytes, bytearray))


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    return str(value)


def _apply_table(anchor: TableAnchor, table: RenderedTable) -> None:
    shape = anchor.shape
    if getattr(shape, "has_table", False):
        _refresh_native_table(anchor, table)
        return
    _replace_placeholder_with_table(anchor, table)


def _refresh_native_table(anchor: TableAnchor, table: RenderedTable) -> None:
    shape = anchor.shape
    native_table = shape.table
    existing_rows = len(native_table.rows)
    existing_columns = len(native_table.columns)
    if existing_rows != table.row_count or existing_columns != table.column_count:
        raise ValueError(
            f"Slide {anchor.slide_number}: table {anchor.token} is {existing_rows}x{existing_columns} "
            f"but data is {table.row_count}x{table.column_count}. "
            f"Adjust the table size or replace it with {anchor.token} and rerun."
        )
    _write_body_cells(native_table, table)
    set_shape_alt_text(shape, anchor.token)


def _replace_placeholder_with_table(anchor: TableAnchor, table: RenderedTable) -> None:
    shape = anchor.shape
    slide = shape.part.slide
    left, top, width, height = shape.left, shape.top, shape.width, shape.height
    default_columns = _default_columns(table)
    parent = shape._element.getparent()
    parent.remove(shape._element)
    table_shape = slide.shapes.add_table(table.row_count, table.column_count, left, top, width, height)
    table_shape.table.first_row = True
    table_shape.table.horz_banding = True
    _write_all_table_cells(table_shape.table, table, default_columns=default_columns)
    set_shape_alt_text(table_shape, anchor.token)


def _write_all_table_cells(native_table: object, table: RenderedTable, *, default_columns: tuple[str, ...]) -> None:
    rows = (default_columns,) + table.body
    for row_index, row in enumerate(rows):
        for column_index, value in enumerate(row):
            native_table.cell(row_index, column_index).text = value


def _write_body_cells(native_table: object, table: RenderedTable) -> None:
    for row_index, row in enumerate(table.body, start=1):
        for column_index, value in enumerate(row):
            native_table.cell(row_index, column_index).text = value
